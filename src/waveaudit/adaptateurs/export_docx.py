"""Adaptateur d'export Word (CAP-6, CAP-8, CAP-13, contrats §17 et §27 [v2], AD-14).

Implémente `domaine.ports.Exporteur`. Le rapport d'incident (CERT, contrats §21) réutilise
`BANDEAU` sans modifier ce module.
"""

import os
import re
import uuid
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.table import Table

from waveaudit.adaptateurs.theme_generique import (
    COULEUR_PRIMAIRE_DEFAUT,
    COULEUR_SECONDAIRE_DEFAUT,
    appliquer_couleurs,
    assurer_styles_titre,
    charger_theme,
    document_generique,
)
from waveaudit.domaine.erreurs import ConfigurationInvalide
from waveaudit.domaine.modeles import CodeMotif, EntreeEcartee, Fiche, Langue, Motif, Rapport
from waveaudit.domaine.ports import OptionsExport

BANDEAU = "Démonstration — données fictives"

# [v2] Libellés par langue (contrats §27, formats.md « Libellés Word par langue »). Les valeurs
# françaises restent celles de contrats §17, inchangées depuis la v0.1.0.
LIBELLES: dict[Langue, dict[str, object]] = {
    "fr": {
        "titre": "Rapport de constats",
        "synthese": ("N°", "Titre", "Sévérité", "Score", "Statut"),
        "description": "Description",
        "impact": "Impact",
        "recommandation": "Recommandation",
        "cvss": "CVSS 3.1",
        "references": "Références",
        "statut": "Statut",
        "tracabilite": "Traçabilité",
        "severites": {
            "aucune": "Aucune",
            "faible": "Faible",
            "moyenne": "Moyenne",
            "elevee": "Élevée",
            "critique": "Critique",
        },
        "ok": "OK",
        "a_revoir": "À revoir",
        "annexe_titre": "Annexe — notes écartées",
        "annexe_entetes": ("Note", "Statut", "Motif"),
        "marque": "Non validée",
        "validee": "Validée",
        "motif_non_validee": "fiche non validée, exclue du rapport final",
        # Mots inline des rubriques CVSS et Traçabilité (contrats §17), hors table §27.
        "vecteur": "Vecteur",
        "score": "Score",
        "severite": "Sévérité",
        "note": "Note",
        "constat_type": "Constat type",
        # [LIB] Libellés des statuts d'entrée écartée (annexe §27) et des codes de motif
        # (contrats §17 « code : detail »), par code de `CodeMotif` (domaine/modeles.py).
        "statuts_ecartes": {
            "non_traitee": "Non traitée",
            "non_rattachee": "Non rattachée",
        },
        "motifs": {
            "reference_retiree": "Référence retirée",
            "vecteur_invalide": "Vecteur CVSS invalide",
            "rubrique_completee": "Rubrique complétée",
            "note_tronquee": "Note tronquée",
            "repli_sortie_invalide": "Repli — sortie invalide",
            "repli_moteur": "Repli moteur",
            "note_invalide": "Note invalide",
            "aucun_constat": "Aucun constat retenu",
            "indicateur_retire": "Indicateur retiré",
            "repetition": "Rubrique répétitive",
            "langue": "Langue incohérente",
            "lot_interrompu": "Lot interrompu",
        },
    },
    "en": {
        "titre": "Findings report",
        "synthese": ("No.", "Title", "Severity", "Score", "Status"),
        "description": "Description",
        "impact": "Impact",
        "recommandation": "Recommendation",
        "cvss": "CVSS 3.1",
        "references": "References",
        "statut": "Status",
        "tracabilite": "Traceability",
        "severites": {
            "aucune": "None",
            "faible": "Low",
            "moyenne": "Medium",
            "elevee": "High",
            "critique": "Critical",
        },
        "ok": "OK",
        "a_revoir": "To review",
        "annexe_titre": "Appendix — excluded notes",
        "annexe_entetes": ("Note", "Status", "Reason"),
        "marque": "Not validated",
        "validee": "Validated",
        "motif_non_validee": "not validated, excluded from the final report",
        "vecteur": "Vector",
        "score": "Score",
        "severite": "Severity",
        "note": "Note",
        "constat_type": "Finding template",
        "statuts_ecartes": {
            "non_traitee": "Not processed",
            "non_rattachee": "Not matched",
        },
        "motifs": {
            "reference_retiree": "Reference removed",
            "vecteur_invalide": "Invalid CVSS vector",
            "rubrique_completee": "Section completed",
            "note_tronquee": "Note truncated",
            "repli_sortie_invalide": "Fallback — invalid output",
            "repli_moteur": "Engine fallback",
            "note_invalide": "Invalid note",
            "aucun_constat": "No matching finding",
            "indicateur_retire": "Indicator removed",
            "repetition": "Repetitive section",
            "langue": "Language mismatch",
            "lot_interrompu": "Batch interrupted",
        },
    },
}

# [LIB] Codes dont `detail` porte une valeur réutilisable (jeton, vecteur CVSS, nom de
# rubrique, message d'erreur technique) : affichée dans les deux langues. Les autres codes
# écrivent une phrase française (contrats §6, §21, §25) : son `detail` n'est repris qu'en
# français ; en anglais, seul le libellé est affiché. `LOT_INTERROMPU` n'a qu'un détail fixe,
# toujours redondant avec son libellé : jamais affiché, dans aucune langue.
_MOTIFS_VALEUR: frozenset[CodeMotif] = frozenset(
    {
        CodeMotif.REFERENCE_RETIREE,
        CodeMotif.VECTEUR_INVALIDE,
        CodeMotif.RUBRIQUE_COMPLETEE,
        CodeMotif.REPETITION,
        CodeMotif.INDICATEUR_RETIRE,
        CodeMotif.REPLI_SORTIE_INVALIDE,
        CodeMotif.REPLI_MOTEUR,
    }
)
_MOTIFS_SANS_DETAIL: frozenset[CodeMotif] = frozenset({CodeMotif.LOT_INTERROMPU})

# [v2] `BANDEAU` reste la constante publique française (contrats §17).
BANDEAUX: dict[Langue, str] = {"fr": BANDEAU, "en": "Demonstration — fictitious data"}

# Caractères interdits en XML 1.0 : python-docx lève ValueError s'il en reçoit un.
_HORS_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff￾￿]")


class ExportDocx:
    """Exporte un `Rapport` en `.docx` (contrats §17, §27). Thème privé s'il est fourni."""

    def __init__(self, theme_prive: Path) -> None:
        """Valide le thème privé dès la composition : `ConfigurationInvalide` avant le lot."""
        self._theme_prive = Path(theme_prive)
        self._theme = charger_theme(self._theme_prive)
        self._modele = self._theme_prive / "modele.docx"
        if self._modele.is_file():
            self._ouvrir_modele()

    def exporter(
        self,
        rapport: Rapport,
        chemin: Path,
        options: OptionsExport = OptionsExport(),  # noqa: B008 - dataclass figée, immuable
    ) -> None:
        libelles = LIBELLES[options.langue]
        document = self._document_de_depart(self._theme)

        _bandeau_en_tete_pied(document, BANDEAUX[options.langue])
        document.add_heading(str(libelles["titre"]), level=0)
        _table_synthese(document, rapport.fiches, libelles, options)
        for index, fiche in enumerate(rapport.fiches, start=1):
            _section_fiche(document, index, fiche, libelles, options)
        if rapport.ecartees or (options.mode == "final" and options.non_validees):
            _annexe(document, rapport.ecartees, options, libelles)

        chemin = Path(chemin)
        chemin.parent.mkdir(parents=True, exist_ok=True)
        # [revue E19] Provisoire unique : deux exports simultanés (service, contrats §27) ne se
        # mélangent pas même si un verrou amont venait à manquer.
        provisoire = chemin.with_name(f".{chemin.name}.{uuid.uuid4().hex}.partiel")
        try:
            document.save(str(provisoire))
            os.replace(provisoire, chemin)
        finally:
            provisoire.unlink(missing_ok=True)

    def _ouvrir_modele(self) -> Document:
        try:
            return Document(str(self._modele))
        except Exception as erreur:  # zip corrompu, pas un .docx : python-docx lève divers types
            raise ConfigurationInvalide(
                f"modele.docx du thème privé illisible : {erreur}"
            ) from erreur

    def _document_de_depart(self, theme: dict) -> Document:
        if self._modele.is_file():
            document = self._ouvrir_modele()
            assurer_styles_titre(document)
        else:
            document = document_generique()
        couleur_primaire = theme.get("couleur_primaire")
        couleur_secondaire = theme.get("couleur_secondaire")
        if couleur_primaire or couleur_secondaire:
            appliquer_couleurs(
                document,
                str(couleur_primaire or COULEUR_PRIMAIRE_DEFAUT),
                str(couleur_secondaire or COULEUR_SECONDAIRE_DEFAUT),
            )
        return document


def _propre(texte: str) -> str:
    """Retire les caractères interdits en XML (un NUL venu du modèle ferait échouer l'export)."""
    return _HORS_XML.sub("", texte)


def _paragraphe(document: Document, texte: str) -> None:
    document.add_paragraph(_propre(texte))


def _conteneurs_bandeau(section, pages_paires: bool):
    yield section.header
    yield section.footer
    if section.different_first_page_header_footer:
        yield section.first_page_header
        yield section.first_page_footer
    if pages_paires:
        yield section.even_page_header
        yield section.even_page_footer


def _bandeau_en_tete_pied(document: Document, bandeau: str = BANDEAU) -> None:
    """Bandeau dans chaque en-tête et pied, sans effacer le contenu d'un modèle (logo, mentions).

    Un en-tête lié au précédent partage sa définition : le test `bandeau in …` évite le doublon.
    """
    pages_paires = document.settings.odd_and_even_pages_header_footer
    for section in document.sections:
        for conteneur in _conteneurs_bandeau(section, pages_paires):
            paragraphes = conteneur.paragraphs
            if any(p.text == bandeau for p in paragraphes):
                continue
            vide = paragraphes and not paragraphes[0].text and not paragraphes[0].runs
            paragraphe = paragraphes[0] if vide else conteneur.add_paragraph()
            paragraphe.text = bandeau
            paragraphe.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _style_table(document: Document, style_prefere: str) -> str | None:
    """`style_prefere` s'il existe (modèle privé), sinon `Table Grid`, sinon le style par défaut."""
    noms = {s.name for s in document.styles}
    for nom in (style_prefere, "Table Grid"):
        if nom in noms:
            return nom
    return None


def _motif_affiche(motif: Motif, libelles: dict, langue: Langue) -> str:
    """[LIB] Libellé du code, avec le détail seulement s'il reste lisible dans `langue`.

    Code inconnu de la table (défensif, `CodeMotif` étant fermé en pratique) : repli sur le
    code brut, sans détail.
    """
    libelle = libelles["motifs"].get(str(motif.code), str(motif.code))
    if motif.code in _MOTIFS_SANS_DETAIL or libelle == str(motif.code):
        return libelle
    if langue == "fr" or motif.code in _MOTIFS_VALEUR:
        return f"{libelle} : {motif.detail}"
    return libelle


def _statut_ecartee_affiche(statut: str, libelles: dict) -> str:
    """[LIB] Libellé du statut d'une entrée écartée ; repli sur le statut brut si inconnu."""
    return libelles["statuts_ecartes"].get(str(statut), str(statut))


def _statut_affiche(fiche: Fiche, libelles: dict, options: OptionsExport) -> str:
    """[v2] En mode final, « Validée », sans égard au statut de génération (contrats §27)."""
    if options.mode == "final":
        return str(libelles["validee"])
    return str(libelles["ok"]) if fiche.statut == "ok" else str(libelles["a_revoir"])


def _table_synthese(
    document: Document, fiches: list[Fiche], libelles: dict, options: OptionsExport
) -> Table:
    table = document.add_table(rows=1, cols=5)
    table.style = _style_table(document, "Light Grid Accent 1")
    entetes = table.rows[0].cells
    for cellule, texte in zip(entetes, libelles["synthese"], strict=True):
        cellule.text = texte
    severites = libelles["severites"]
    for index, fiche in enumerate(fiches, start=1):
        cellules = table.add_row().cells
        valeurs = (
            str(index),
            fiche.titre,
            severites[fiche.cvss.severite],
            f"{fiche.cvss.score:.1f}",
            _statut_affiche(fiche, libelles, options),
        )
        for cellule, texte in zip(cellules, valeurs, strict=True):
            cellule.text = _propre(texte)
    return table


def _marque_titre(fiche: Fiche, libelles: dict, options: OptionsExport) -> str:
    """[v2] « — Non validée » dans le titre, en brouillon seulement (contrats §27)."""
    if options.mode == "brouillon" and fiche.tracabilite.note in options.non_validees:
        return f" — {libelles['marque']}"
    return ""


def _section_fiche(
    document: Document, index: int, fiche: Fiche, libelles: dict, options: OptionsExport
) -> None:
    marque = _marque_titre(fiche, libelles, options)
    document.add_heading(_propre(f"{index}. {fiche.titre}{marque}"), level=2)

    document.add_heading(str(libelles["description"]), level=3)
    _paragraphe(document, fiche.description)

    document.add_heading(str(libelles["impact"]), level=3)
    _paragraphe(document, fiche.impact)

    document.add_heading(str(libelles["recommandation"]), level=3)
    _paragraphe(document, fiche.recommandation)

    document.add_heading(str(libelles["cvss"]), level=3)
    severites = libelles["severites"]
    _paragraphe(
        document,
        f"{libelles['vecteur']} : {fiche.cvss.vecteur} — "
        f"{libelles['score']} : {fiche.cvss.score:.1f} — "
        f"{libelles['severite']} : {severites[fiche.cvss.severite]}",
    )

    document.add_heading(str(libelles["references"]), level=3)
    for reference in fiche.references:
        _paragraphe(document, f"{reference.id} — {reference.titre} — {reference.url}")

    document.add_heading(str(libelles["statut"]), level=3)
    if options.mode == "final":
        document.add_paragraph(str(libelles["validee"]))
    elif fiche.statut == "ok":
        document.add_paragraph(str(libelles["ok"]))
    else:
        document.add_paragraph(str(libelles["a_revoir"]))
        for motif in fiche.motifs:
            _paragraphe(document, _motif_affiche(motif, libelles, options.langue))

    document.add_heading(str(libelles["tracabilite"]), level=3)
    _paragraphe(
        document,
        f"{libelles['note']} : {fiche.tracabilite.note} — "
        f"{libelles['constat_type']} : {fiche.tracabilite.constat_id}",
    )


def _annexe(
    document: Document, ecartees: list[EntreeEcartee], options: OptionsExport, libelles: dict
) -> None:
    document.add_heading(str(libelles["annexe_titre"]), level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = _style_table(document, "Table Grid")
    entetes = table.rows[0].cells
    for cellule, texte in zip(entetes, libelles["annexe_entetes"], strict=True):
        cellule.text = texte
    for entree in ecartees:
        cellules = table.add_row().cells
        valeurs = (
            entree.note,
            _statut_ecartee_affiche(entree.statut, libelles),
            _motif_affiche(entree.motif, libelles, options.langue),
        )
        for cellule, texte in zip(cellules, valeurs, strict=True):
            cellule.text = _propre(texte)
    if options.mode == "final":
        # [revue A17] Les fiches non validées sont listées en annexe, jamais dans le corps.
        for note in sorted(options.non_validees, key=str.casefold):
            cellules = table.add_row().cells
            valeurs = (note, str(libelles["marque"]), str(libelles["motif_non_validee"]))
            for cellule, texte in zip(cellules, valeurs, strict=True):
                cellule.text = _propre(texte)
