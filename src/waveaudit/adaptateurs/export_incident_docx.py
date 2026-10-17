"""Adaptateur d'export Word du rapport d'incident (CAP-6, contrats §21, AD-17).

Implémente `domaine.ports.ExporteurIncident`. Sections, intitulés et mise en forme :
`formats-incident.md` « Rapport d'incident ». Réutilise le bandeau et le nettoyage XML
d'`export_docx.py`, et le thème générique de `theme_generique.py`, sans modifier ces modules.
"""

import os
from pathlib import Path

from docx import Document
from docx.table import Table

from waveaudit.adaptateurs.export_docx import _bandeau_en_tete_pied, _propre, _style_table
from waveaudit.adaptateurs.theme_generique import (
    COULEUR_PRIMAIRE_DEFAUT,
    COULEUR_SECONDAIRE_DEFAUT,
    appliquer_couleurs,
    assurer_styles_titre,
    charger_theme,
    document_generique,
)
from waveaudit.domaine.erreurs import ConfigurationInvalide
from waveaudit.domaine.incident import (
    LIBELLES_SEVERITE,
    SEVERITES,
    Indicateur,
    RapportIncident,
    Recit,
)
from waveaudit.domaine.modeles import Statut

_AUCUN_FAIT = "Aucun fait détecté"


class ExporteurIncidentDocx:
    """Exporte un `RapportIncident` en `.docx` (contrats §21). Thème privé si fourni."""

    def __init__(self, theme_prive: Path, max_preuves: int = 10) -> None:
        """Valide le thème privé dès la composition : `ConfigurationInvalide` avant le lot."""
        if max_preuves < 1:  # même borne que `config.incident.max_preuves` (ge=1)
            raise ConfigurationInvalide(f"max_preuves doit valoir au moins 1 : {max_preuves}")
        self._theme_prive = Path(theme_prive)
        self._max_preuves = max_preuves
        self._theme = charger_theme(self._theme_prive)
        self._modele = self._theme_prive / "modele.docx"
        if self._modele.is_file():
            self._ouvrir_modele()

    def exporter(self, rapport: RapportIncident, chemin: Path) -> None:
        document = self._document_de_depart(self._theme)

        _bandeau_en_tete_pied(document)
        document.add_heading("Rapport d'incident", level=0)
        _section_synthese(document, rapport)
        _section_chronologie(document, rapport.recits)
        _section_faits(document, rapport.recits, self._max_preuves)
        _section_indicateurs(document, rapport.indicateurs)
        _section_actifs(document, rapport.actifs)
        _section_recommandations(document, rapport.recommandations)
        _section_sources(document, rapport)

        chemin = Path(chemin)
        chemin.parent.mkdir(parents=True, exist_ok=True)
        provisoire = chemin.with_name(f".{chemin.name}.partiel")
        try:
            document.save(str(provisoire))
            os.replace(provisoire, chemin)
        finally:
            provisoire.unlink(missing_ok=True)

    def _ouvrir_modele(self) -> Document:
        try:
            return Document(str(self._modele))
        except Exception as erreur:  # zip corrompu, pas un .docx : python-docx lève divers types
            raise ConfigurationInvalide(
                f"modele.docx du thème privé illisible : {erreur}"
            ) from erreur

    def _document_de_depart(self, theme: dict) -> Document:
        if self._modele.is_file():
            document = self._ouvrir_modele()
            assurer_styles_titre(document)
        else:
            document = document_generique()
        couleur_primaire = theme.get("couleur_primaire")
        couleur_secondaire = theme.get("couleur_secondaire")
        if couleur_primaire or couleur_secondaire:
            appliquer_couleurs(
                document,
                str(couleur_primaire or COULEUR_PRIMAIRE_DEFAUT),
                str(couleur_secondaire or COULEUR_SECONDAIRE_DEFAUT),
            )
        return document


def _paragraphe(document: Document, texte: str) -> None:
    document.add_paragraph(_propre(texte))


def _section_synthese(document: Document, rapport: RapportIncident) -> None:
    document.add_heading("Synthèse", level=1)
    if rapport.debut is not None and rapport.fin is not None:
        _paragraphe(
            document,
            f"Période : {rapport.debut.isoformat(timespec='seconds')} → "
            f"{rapport.fin.isoformat(timespec='seconds')}",
        )
    else:
        _paragraphe(document, "Période : aucune")

    table = document.add_table(rows=1, cols=2)
    table.style = _style_table(document, "Light Grid Accent 1")
    for cellule, texte in zip(table.rows[0].cells, ("Sévérité", "Nombre"), strict=True):
        cellule.text = texte
    comptes = {severite: 0 for severite in SEVERITES}
    for recit in rapport.recits:
        comptes[recit.severite] += 1
    for severite in SEVERITES:
        cellules = table.add_row().cells
        for cellule, texte in zip(
            cellules, (LIBELLES_SEVERITE[severite], str(comptes[severite])), strict=True
        ):
            cellule.text = _propre(texte)

    if rapport.actifs:
        _paragraphe(document, "Hôtes touchés : " + ", ".join(rapport.actifs))
    else:
        _paragraphe(document, "Hôtes touchés : aucun")


def _section_chronologie(document: Document, recits: list[Recit]) -> None:
    document.add_heading("Chronologie", level=1)
    if not recits:
        document.add_paragraph(_AUCUN_FAIT)
        return
    table = document.add_table(rows=1, cols=4)
    table.style = _style_table(document, "Table Grid")
    entetes = ("Horodatage", "Hôte", "Fait", "Sévérité")
    for cellule, texte in zip(table.rows[0].cells, entetes, strict=True):
        cellule.text = texte
    for recit in recits:
        detection = recit.detection
        cellules = table.add_row().cells
        valeurs = (
            detection.debut.isoformat(timespec="seconds"),
            ", ".join(detection.hotes),
            recit.regle_titre,
            LIBELLES_SEVERITE[recit.severite],
        )
        for cellule, texte in zip(cellules, valeurs, strict=True):
            cellule.text = _propre(texte)


def _section_faits(document: Document, recits: list[Recit], max_preuves: int) -> None:
    document.add_heading("Faits détectés", level=1)
    if not recits:
        document.add_paragraph(_AUCUN_FAIT)
        return
    for index, recit in enumerate(recits, start=1):
        _section_recit(document, index, recit, max_preuves)


def _section_recit(document: Document, index: int, recit: Recit, max_preuves: int) -> None:
    document.add_heading(_propre(f"{index}. {recit.titre} ({recit.detection.regle_id})"), level=2)

    document.add_heading("Récit", level=3)
    _paragraphe(document, recit.recit)

    document.add_heading("Impact", level=3)
    _paragraphe(document, recit.impact)

    document.add_heading("Recommandation", level=3)
    _paragraphe(document, recit.recommandation)

    document.add_heading("Références ATT&CK", level=3)
    for reference in recit.references:
        _paragraphe(document, f"{reference.id} — {reference.titre} — {reference.url}")

    document.add_heading("Preuves", level=3)
    evenements = recit.detection.evenements
    for evenement in evenements[:max_preuves]:
        message = " ".join(evenement.message.split())[:300]
        _paragraphe(document, f"{evenement.fichier}:{evenement.ligne} — {message}")
    reste = len(evenements) - max_preuves
    if reste > 0:
        _paragraphe(document, f"+ {reste} autres")

    document.add_heading("Statut", level=3)
    if recit.statut == Statut.OK:
        document.add_paragraph("OK")
    else:
        document.add_paragraph("À revoir")
        for motif in recit.motifs:
            _paragraphe(document, f"{motif.code} : {motif.detail}")


def _section_indicateurs(document: Document, indicateurs: list[Indicateur]) -> None:
    document.add_heading("Indicateurs de compromission", level=1)
    if not indicateurs:
        document.add_paragraph(_AUCUN_FAIT)
        return
    table = document.add_table(rows=1, cols=3)
    table.style = _style_table(document, "Table Grid")
    for cellule, texte in zip(table.rows[0].cells, ("Type", "Valeur", "Origines"), strict=True):
        cellule.text = texte
    for indicateur in indicateurs:
        cellules = table.add_row().cells
        valeurs = (indicateur.type, indicateur.valeur, ", ".join(indicateur.origines))
        for cellule, texte in zip(cellules, valeurs, strict=True):
            cellule.text = _propre(texte)


def _section_actifs(document: Document, actifs: list[str]) -> None:
    document.add_heading("Actifs touchés", level=1)
    if not actifs:
        document.add_paragraph(_AUCUN_FAIT)
        return
    for actif in actifs:
        _paragraphe(document, actif)


def _section_recommandations(document: Document, recommandations: list[str]) -> None:
    document.add_heading("Recommandations consolidées", level=1)
    if not recommandations:
        document.add_paragraph(_AUCUN_FAIT)
        return
    for recommandation in recommandations:
        _paragraphe(document, recommandation)


def _section_sources(document: Document, rapport: RapportIncident) -> Table:
    document.add_heading("Annexe : sources analysées", level=1)
    ignorees_par_fichier: dict[str, int] = {}
    for ligne in rapport.ignorees:
        ignorees_par_fichier[ligne.fichier] = ignorees_par_fichier.get(ligne.fichier, 0) + 1

    table = document.add_table(rows=1, cols=3)
    table.style = _style_table(document, "Table Grid")
    entetes = ("Fichier", "Enregistrements examinés", "Ignorés")
    for cellule, texte in zip(table.rows[0].cells, entetes, strict=True):
        cellule.text = texte
    for fichier in sorted(rapport.sources, key=str.casefold):
        cellules = table.add_row().cells
        valeurs = (
            fichier,
            str(rapport.sources[fichier]),
            str(ignorees_par_fichier.get(fichier, 0)),
        )
        for cellule, texte in zip(cellules, valeurs, strict=True):
            cellule.text = _propre(texte)
    for fichier in sorted(rapport.sources, key=str.casefold):
        if fichier.casefold().endswith(".json"):
            _paragraphe(document, f"{fichier} : ligne = rang d'objet")
    return table
