"""Service incident de bout en bout (contrats §21 « Service », CAP-3 à CAP-5), faux moteur.

L'oracle `donnees/demo/incident-attendu.yaml` est écrit à la main : ces tests le comparent à
la sortie du service, jamais l'inverse (formats-incident.md « Jeu de démonstration et oracle »).
"""

import json
import time
from datetime import datetime
from pathlib import Path

import pytest
import yaml
from docx import Document

from waveaudit.application import composition
from waveaudit.application.config import racine_projet
from waveaudit.application.incident import executer_incident
from waveaudit.domaine.erreurs import EntreeInvalide, MoteurInjoignable
from waveaudit.domaine.incident import RapportIncident, RequeteIncident
from waveaudit.domaine.modeles import CodeMotif, Statut, TypeEvenement

RACINE = racine_projet()
DEMO = RACINE / "donnees" / "demo" / "incident"
ORACLE = yaml.safe_load(
    (RACINE / "donnees" / "demo" / "incident-attendu.yaml").read_text(encoding="utf-8")
)
ANNEE = 2026
T = TypeEvenement


def _vers(config, fm):
    return config.model_copy(
        update={"moteur": config.moteur.model_copy(update={"base_url": fm.base_url})}
    )


def _journal(resultat_ou_dossier) -> list[dict]:
    dossier = getattr(resultat_ou_dossier, "dossier_sortie", resultat_ou_dossier)
    texte = (Path(dossier) / "journal.jsonl").read_text(encoding="utf-8")
    return [json.loads(ligne) for ligne in texte.splitlines()]


def _executer(config, dossier=DEMO, **requete):
    return executer_incident(RequeteIncident(dossier=dossier, annee=ANNEE, **requete), config)


def _codes(recit) -> list[CodeMotif]:
    return [m.code for m in recit.motifs]


def test_demo_egale_a_l_oracle(config_test, tmp_path):
    assert ORACLE["annee"] == ANNEE
    assert ORACLE["decalage_syslog"] == config_test.incident.decalage_syslog
    recus = []
    copie = tmp_path / "copie.docx"
    chrono = time.perf_counter()
    resultat = executer_incident(
        RequeteIncident(dossier=DEMO, annee=ANNEE, copie_rapport=copie),
        config_test,
        recus.append,
    )
    assert time.perf_counter() - chrono < 30
    rapport = resultat.rapport

    detections = [r.detection for r in rapport.recits]
    assert [(d.regle_id, d.debut, d.fin, d.hotes, len(d.evenements)) for d in detections] == [
        (
            a["regle_id"],
            datetime.fromisoformat(a["debut"]),
            datetime.fromisoformat(a["fin"]),
            a["hotes"],
            a["nb_evenements"],
        )
        for a in ORACLE["detections"]
    ]
    assert len(detections) == 16
    assert [i.model_dump() for i in rapport.indicateurs] == ORACLE["indicateurs"]
    assert [(i.fichier, i.ligne) for i in rapport.ignorees] == [
        (i["fichier"], i["ligne"]) for i in ORACLE["ignorees"]
    ]

    lus = sum(rapport.sources.values()) - len(rapport.ignorees)
    detectes = {e.cle for d in detections for e in d.evenements}
    assert lus >= 80
    assert (lus - len(detectes)) / lus >= ORACLE["min_part_benigne"]

    assert all(r.statut == Statut.OK for r in rapport.recits)
    assert all(r.recit.startswith("[adapté] ") for r in rapport.recits)
    assert rapport.debut == detections[0].debut and rapport.fin == max(d.fin for d in detections)

    sortie = resultat.dossier_sortie
    assert sortie == config_test.chemins.var / "sorties" / resultat.id_lot
    json_brut = (sortie / "incident.json").read_text(encoding="utf-8")
    assert RapportIncident.model_validate_json(json_brut) == rapport
    assert json_brut == rapport.model_dump_json(indent=2)
    assert resultat.chemin_rapport == sortie / "rapport-incident.docx"
    Document(str(resultat.chemin_rapport))
    assert copie.read_bytes() == resultat.chemin_rapport.read_bytes()

    evenements = _journal(resultat)
    assert [e["type"] for e in evenements] == [e.type for e in recus]
    assert evenements[0]["type"] == T.LOT_DEMARRE and evenements[-1]["type"] == T.LOT_TERMINE
    assert evenements[0]["total"] == 16
    assert evenements[0]["details"] == {"workflow": "incident", "pid_moteur": None}
    paires = evenements[1:-1]
    assert len(paires) == 32
    for i, d in enumerate(detections):
        debut, fin = paires[2 * i], paires[2 * i + 1]
        assert (debut["type"], debut["index"], debut["note"]) == (
            T.FICHE_DEMARREE,
            i + 1,
            d.regle_id,
        )
        assert (fin["type"], fin["index"], fin["note"]) == (T.FICHE_TERMINEE, i + 1, d.regle_id)
        assert fin["statut"] == Statut.OK and fin["duree_s"] >= 0
    assert str(tmp_path) not in (sortie / "journal.jsonl").read_text(encoding="utf-8")


def test_deux_executions_incident_json_identique(config_test):
    premier = _executer(config_test)
    second = _executer(config_test)
    assert premier.id_lot != second.id_lot
    assert (premier.dossier_sortie / "incident.json").read_bytes() == (
        second.dossier_sortie / "incident.json"
    ).read_bytes()


def test_json_invalide_une_fois(config_test, faux_moteur):
    fm = faux_moteur("json_invalide_une_fois")
    resultat = _executer(_vers(config_test, fm))
    assert len(fm.requetes) == 17
    assert all(r.statut == Statut.OK for r in resultat.rapport.recits)


def test_erreur_500_repli_a_revoir(config_test, faux_moteur):
    fm = faux_moteur("erreur_500")
    resultat = _executer(_vers(config_test, fm))
    recits = resultat.rapport.recits
    assert len(recits) == 16
    for r in recits:
        assert r.statut == Statut.A_REVOIR and _codes(r) == [CodeMotif.REPLI_MOTEUR]
    fins = [e for e in _journal(resultat) if e["type"] == T.FICHE_TERMINEE]
    assert [e["statut"] for e in fins] == [Statut.A_REVOIR] * 16


def test_references_inventees_retirees(config_test, faux_moteur):
    fm = faux_moteur("references_inventees")
    resultat = _executer(_vers(config_test, fm))
    attendus = sorted([CodeMotif.REFERENCE_RETIREE] * 2 + [CodeMotif.INDICATEUR_RETIRE] * 2)
    for r in resultat.rapport.recits:
        assert r.statut == Statut.A_REVOIR
        assert sorted(_codes(r)) == attendus
    rubriques = " ".join(r.recit for r in resultat.rapport.recits)
    for jeton in ("T9999", "T1110.999", "203.0.113.254", "0f" * 32):
        assert jeton not in rubriques
    assert "203.0.113.254" not in [i.valeur for i in resultat.rapport.indicateurs]


def test_sans_detection_moteur_injoignable_code_0(config_test, faux_moteur, tmp_path, monkeypatch):
    def interdit(*args, **kwargs):
        raise AssertionError("le moteur ne doit pas être ouvert sans détection")

    monkeypatch.setattr(composition, "ouvrir_moteur_incident", interdit)
    dossier = tmp_path / "journaux"
    dossier.mkdir()
    (dossier / "auth.log").write_text(
        "Oct 14 08:00:01 BASTION01 sshd[100]: Server listening on 0.0.0.0 port 22.\n",
        encoding="utf-8",
    )
    resultat = _executer(_vers(config_test, faux_moteur("injoignable")), dossier)
    rapport = resultat.rapport
    assert rapport.recits == [] and rapport.debut is None and rapport.indicateurs == []
    assert rapport.sources == {"auth.log": 1}
    texte = "\n".join(p.text for p in Document(str(resultat.chemin_rapport)).paragraphs)
    assert "Aucun fait détecté" in texte
    evenements = _journal(resultat)
    assert [e["type"] for e in evenements] == [T.LOT_DEMARRE, T.LOT_TERMINE]
    assert evenements[0]["total"] == 0 and evenements[0]["details"] == {"workflow": "incident"}


def test_demo_moteur_injoignable(config_test, faux_moteur):
    config = _vers(config_test, faux_moteur("injoignable"))
    with pytest.raises(MoteurInjoignable):
        _executer(config)
    (dossier,) = (config.chemins.var / "sorties").iterdir()
    evenements = _journal(dossier)
    assert [e["type"] for e in evenements] == [T.LOT_ECHOUE]
    assert evenements[0]["details"]["erreur"] == "MoteurInjoignable"
    assert not (dossier / "incident.json").exists()


def test_dossier_absent(config_test, tmp_path):
    with pytest.raises(EntreeInvalide):
        _executer(config_test, tmp_path / "absent")


def test_dossier_sans_fichier_route(config_test, tmp_path):
    dossier = tmp_path / "vide"
    dossier.mkdir()
    with pytest.raises(EntreeInvalide):
        _executer(config_test, dossier)
    (sortie,) = (config_test.chemins.var / "sorties").iterdir()
    assert [e["type"] for e in _journal(sortie)] == [T.LOT_ECHOUE]


# Revue C5, S0 (v2) : `application.incident` réutilise les noms publics de `lot.py`
# (`MOTIF_ID_LOT`, `Emetteur`, `sans_chemin`, `dossier_sorties`, `verifier_id_lot`) ;
# ces tests cassent si leur comportement change.
@pytest.mark.parametrize("id_lot", ["../evasion", "-tiret", "a b", "NUL"])
def test_id_lot_refuse(config_test, id_lot):
    with pytest.raises(EntreeInvalide):
        _executer(config_test, id_lot=id_lot)


def test_id_lot_fourni_et_sans_fichier_provisoire(config_test, tmp_path):
    copie = tmp_path / "copie.docx"
    resultat = _executer(config_test, id_lot="lot-revue-1", copie_rapport=copie)
    assert resultat.id_lot == "lot-revue-1"
    assert resultat.dossier_sortie == config_test.chemins.var / "sorties" / "lot-revue-1"
    assert sorted(p.name for p in resultat.dossier_sortie.iterdir()) == [
        "incident.json",
        "journal.jsonl",
        "rapport-incident.docx",
    ]
    assert [p.name for p in tmp_path.iterdir() if p.name.startswith(".")] == []
    assert copie.read_bytes() == resultat.chemin_rapport.read_bytes()


def test_lot_echoue_sans_chemin_absolu(config_test, tmp_path):
    dossier = tmp_path / "vide"
    dossier.mkdir()
    with pytest.raises(EntreeInvalide):
        _executer(config_test, dossier)
    (sortie,) = (config_test.chemins.var / "sorties").iterdir()
    (echec,) = _journal(sortie)
    assert echec["details"]["erreur"] == "EntreeInvalide"
    assert str(tmp_path) not in echec["details"]["message"]
