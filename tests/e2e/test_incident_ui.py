"""Workflow incident dans l'UI, de bout en bout (contrats §26, §30, CAP-14).

`client_web(config_test)` est en mode `externe` vers un faux moteur nominal (`tests/conftest.py`) :
l'année courante vaut 2026 dans l'environnement de test, mais chaque test passe `annee=2026`
explicitement (comme la CLI, `tests/e2e/test_cli_incident.py`) pour ne jamais dépendre de la
date d'exécution.
"""

import io
import time
from typing import Any

from docx import Document
from fastapi.testclient import TestClient
from lxml import html as lxml_html

from tests.accessibilite import verifier_accessibilite, verifier_focus_css
from tests.e2e.test_cli import _ecrire_config, runner
from tests.e2e.test_cli import app as cli_app
from tests.web_outils import application
from tests.web_outils import client as client_web
from waveaudit.application.config import racine_projet
from waveaudit.securite import garde_reseau

DEMO = racine_projet() / "donnees" / "demo" / "incident"
ANNEE = "2026"
_RACINE_WEB = racine_projet() / "src" / "waveaudit" / "interfaces" / "web"


def _vers(config, fm):
    return config.model_copy(
        update={"moteur": config.moteur.model_copy(update={"base_url": fm.base_url})}
    )


def _journaux_demo() -> list[tuple[str, tuple[str, bytes]]]:
    return [
        ("journaux", (chemin.name, chemin.read_bytes()))
        for chemin in sorted(DEMO.iterdir())
        if chemin.is_file() and not chemin.name.startswith(".")
    ]


def _journal_sans_detection() -> list[tuple[str, tuple[str, bytes]]]:
    return [
        (
            "journaux",
            (
                "auth.log",
                b"Oct 14 08:00:01 BASTION01 sshd[100]: Server listening on 0.0.0.0 port 22.\n",
            ),
        )
    ]


def _attendre_fin(client: TestClient, id_lot: str, timeout_s: float = 10.0) -> dict:
    limite = time.monotonic() + timeout_s
    while True:
        reponse = client.get(f"/incidents/{id_lot}/etat")
        assert reponse.status_code == 200
        donnees = reponse.json()
        if donnees["etat"] != "en_cours":
            return donnees
        if time.monotonic() > limite:
            raise AssertionError(f"incident {id_lot} toujours en cours après {timeout_s}s")
        time.sleep(0.05)


def _texte_docx(contenu: bytes) -> list[str]:
    """Texte complet d'un `.docx` (paragraphes puis tableaux, dans l'ordre), sans métadonnées
    (horodatage de création) : seul le contenu visible compte pour la comparaison."""
    document = Document(io.BytesIO(contenu))
    morceaux = [p.text for p in document.paragraphs]
    for table in document.tables:
        for ligne in table.rows:
            morceaux.append(" | ".join(cellule.text for cellule in ligne.cells))
    return morceaux


def test_incident_nominal_16_detections_meme_contenu_que_la_cli(config_test, tmp_path):
    # Oracle : la commande CLI, sur le même dépôt et le même faux moteur nominal.
    config_cli = _ecrire_config(tmp_path, base_url=config_test.moteur.base_url)
    copie = tmp_path / "oracle.docx"
    resultat_cli = runner.invoke(
        cli_app,
        ["incident", str(DEMO), "--annee", ANNEE, "--config", str(config_cli), "-o", str(copie)],
    )
    assert resultat_cli.exit_code == 0, resultat_cli.output

    client = client_web(config_test)
    reponse = client.post(
        "/incidents", files=_journaux_demo(), data={"annee": ANNEE}, follow_redirects=False
    )
    assert reponse.status_code == 303
    id_lot = reponse.headers["location"].rsplit("/", 1)[-1]
    assert reponse.headers["location"] == f"/incidents/{id_lot}"

    # [revue W2, consigne W4] `lot.json` écrit avant le 303 : pas de redirection prématurée.
    page = client.get(f"/incidents/{id_lot}", follow_redirects=False)
    assert page.status_code == 200
    verifier_accessibilite(page.text)

    etat = _attendre_fin(client, id_lot)
    assert etat["etat"] == "termine"

    resultat = client.get(f"/incidents/{id_lot}/resultat")
    assert resultat.status_code == 200
    corps = resultat.json()
    assert corps["total"] == 16
    assert len(corps["rapport"]["recits"]) == 16
    assert sum(corps["par_severite"].values()) == 16
    assert set(corps["par_severite"]) == {"critique", "elevee", "moyenne", "faible"}

    docx = client.get(f"/incidents/{id_lot}/rapport-incident.docx")
    assert docx.status_code == 200
    assert docx.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert _texte_docx(docx.content) == _texte_docx(copie.read_bytes())


def test_incident_sans_detection_total_zero_et_docx_telechargeable(config_test):
    client = client_web(config_test)
    reponse = client.post(
        "/incidents",
        files=_journal_sans_detection(),
        data={"annee": ANNEE},
        follow_redirects=False,
    )
    assert reponse.status_code == 303
    id_lot = reponse.headers["location"].rsplit("/", 1)[-1]
    _attendre_fin(client, id_lot)

    resultat = client.get(f"/incidents/{id_lot}/resultat")
    assert resultat.status_code == 200
    corps = resultat.json()
    assert corps["total"] == 0
    assert corps["rapport"]["recits"] == []

    # Message statique du gabarit (EXPERIENCE « Résultat incident, vide ») : affiché par le JS
    # quand `total == 0`, toujours présent dans le HTML rendu.
    page = client.get(f"/incidents/{id_lot}")
    assert "Aucune détection sur ce dépôt" in page.text

    docx = client.get(f"/incidents/{id_lot}/rapport-incident.docx")
    assert docx.status_code == 200


def test_aucun_journal_reconnu_sans_lot_cree(config_test):
    client = client_web(config_test)
    entrees_avant = list(config_test.chemins.var.glob("entrees/*"))
    sorties_avant = list(config_test.chemins.var.glob("sorties/*"))

    reponse = client.post(
        "/incidents",
        files=[("journaux", ("notes.md", b"pas un journal"))],
        follow_redirects=False,
    )

    assert reponse.status_code == 400
    assert reponse.json()["code"] == "aucun_journal_reconnu"
    assert list(config_test.chemins.var.glob("entrees/*")) == entrees_avant
    assert list(config_test.chemins.var.glob("sorties/*")) == sorties_avant


def test_choix_du_modele_comme_pour_les_fiches(config_test):
    client = client_web(config_test)
    modeles = client.get("/api/modeles").json()["modeles"]
    (modele_externe,) = modeles  # mode externe : une seule entrée (contrats §11, §26)
    assert modele_externe["role"] == "externe"

    reponse = client.post(
        "/incidents",
        files=_journaux_demo(),
        data={"annee": ANNEE, "modele": modele_externe["nom"]},
        follow_redirects=False,
    )
    assert reponse.status_code == 303
    id_lot = reponse.headers["location"].rsplit("/", 1)[-1]
    _attendre_fin(client, id_lot)

    fiche_lot = (config_test.chemins.var / "sorties" / id_lot / "lot.json").read_text(
        encoding="utf-8"
    )
    # [revue E11] Le modèle `externe` affiché n'est jamais repassé tel quel au lot.
    assert '"modele_demande": null' in fiche_lot

    refus = client.post(
        "/incidents",
        files=_journaux_demo(),
        data={"annee": ANNEE, "modele": "inconnu-xyz"},
        follow_redirects=False,
    )
    assert refus.status_code == 422
    assert refus.json()["code"] == "modele_indisponible"


def test_verrou_partage_avec_un_lot_de_fiches(config_test, faux_moteur):
    from tests.e2e.test_ui import _notes_demo

    lent = _vers(config_test, faux_moteur("lent"))
    client = client_web(lent)
    premier = client.post("/lots", files=_notes_demo()[:1], follow_redirects=False)
    assert premier.status_code == 303
    id_lot_fiches = premier.headers["location"].rsplit("/", 1)[-1]

    second = client.post("/incidents", files=_journaux_demo(), follow_redirects=False)
    assert second.status_code == 409
    corps = second.json()
    assert corps["code"] == "lot_en_cours"
    assert corps["occupation"]["nature"] == "fiches"
    assert corps["occupation"]["id_lot"] == id_lot_fiches

    from tests.e2e.test_ui import _attendre_fin as _attendre_fin_fiches

    _attendre_fin_fiches(client, id_lot_fiches)


def test_limites_de_depot(config_test):
    client = client_web(config_test)
    entrees = config_test.chemins.var / "entrees"

    # Nom terminé par un point ou une espace : refusé (§26 [revue A6/E41]).
    for nom in ("auth.log.", "auth.log "):
        reponse = client.post(
            "/incidents", files=[("journaux", (nom, b"contenu"))], follow_redirects=False
        )
        assert reponse.status_code == 400, nom
        assert reponse.json()["code"] == "nom_refuse", nom
        assert not entrees.is_dir() or not list(entrees.iterdir())

    # Plus de 50 fichiers : refusé avant toute écriture.
    trop = [("journaux", (f"auth{i}.log", b"x")) for i in range(51)]
    reponse = client.post("/incidents", files=trop, follow_redirects=False)
    assert reponse.status_code == 400
    assert reponse.json()["code"] == "trop_de_fichiers"
    assert not entrees.is_dir() or not list(entrees.iterdir())

    # Plus de 10 Mio pour un seul fichier : refusé (413, limite "fichier").
    enorme = [("journaux", ("auth.log", b"a" * (10 * 1024 * 1024 + 1)))]
    reponse = client.post("/incidents", files=enorme, follow_redirects=False)
    assert reponse.status_code == 413
    assert reponse.json()["limite"] == "fichier"
    assert not entrees.is_dir() or not list(entrees.iterdir())

    # Un envoi valide reste accepté après ces refus (verrou bien relâché).
    reponse = client.post(
        "/incidents", files=_journaux_demo(), data={"annee": ANNEE}, follow_redirects=False
    )
    assert reponse.status_code == 303
    id_lot = reponse.headers["location"].rsplit("/", 1)[-1]
    _attendre_fin(client, id_lot)


def test_annee_invalide(config_test):
    client = client_web(config_test)
    # [revue W4] "--5" et 5 000 chiffres faisaient lever `int()` (erreur 500).
    for annee in ("1900", "2200", "pas-une-annee", "--5", "9" * 5000, "٢٠٢٦"):
        reponse = client.post(
            "/incidents",
            files=_journaux_demo(),
            data={"annee": annee},
            follow_redirects=False,
        )
        assert reponse.status_code == 422, annee
        assert reponse.json()["code"] == "annee_invalide", annee


def test_page_d_un_lot_de_fiches_redirige_vers_lots(config_test):
    from tests.e2e.test_ui import _notes_demo

    client = client_web(config_test)
    reponse = client.post("/lots", files=_notes_demo()[:1], follow_redirects=False)
    id_lot = reponse.headers["location"].rsplit("/", 1)[-1]

    redirige = client.get(f"/incidents/{id_lot}", follow_redirects=False)
    assert redirige.status_code == 303
    assert redirige.headers["location"] == f"/lots/{id_lot}"

    from tests.e2e.test_ui import _attendre_fin as _attendre_fin_fiches

    _attendre_fin_fiches(client, id_lot)


def test_accessibilite_incident(config_test):
    client = client_web(config_test)
    reponse = client.post(
        "/incidents", files=_journaux_demo(), data={"annee": ANNEE}, follow_redirects=False
    )
    id_lot = reponse.headers["location"].rsplit("/", 1)[-1]
    page = client.get(f"/incidents/{id_lot}")
    assert page.status_code == 200
    verifier_accessibilite(page.text)
    # Groupe de boutons bascule (filtre par sévérité, EXPERIENCE « Page incident ») : rendu
    # statique dans le gabarit (DESIGN.md), un bouton par sévérité, non pressé au départ.
    for severite in ("critique", "elevee", "moyenne", "faible"):
        assert f'data-severite="{severite}" aria-pressed="false"' in page.text
    verifier_focus_css((_RACINE_WEB / "static" / "style.css").read_text(encoding="utf-8"))

    _attendre_fin(client, id_lot)


def test_ui_incident_sous_garde_sans_tentative(monkeypatch, tmp_path, faux_moteur):
    """Comme `e2e/test_hors_ligne.py::test_ui_lot_sous_garde_sans_tentative`, pour l'incident :
    aucune tentative de connexion hors boucle locale (CAP-11, contrats §12)."""
    config = _ecrire_config(tmp_path, base_url=faux_moteur().base_url)
    capture: dict[str, Any] = {}

    def faux_run(application: Any, *, host: str, port: int) -> None:
        capture["garde_actif_au_demarrage"] = garde_reseau._actif
        client = client_web(application, base_url=f"http://127.0.0.1:{port}")
        reponse = client.post(
            "/incidents", files=_journaux_demo(), data={"annee": ANNEE}, follow_redirects=False
        )
        id_lot = reponse.headers["location"].rsplit("/", 1)[-1]
        capture["etat"] = _attendre_fin(client, id_lot)

    monkeypatch.setattr("uvicorn.run", faux_run)
    resultat = runner.invoke(cli_app, ["ui", "--config", str(config)])

    assert resultat.exit_code == 0, resultat.output
    assert capture["garde_actif_au_demarrage"]
    assert capture["etat"]["etat"] == "termine"
    assert garde_reseau.tentatives() == []
    assert not garde_reseau._actif


def test_accueil_formulaire_incident_aboutit_au_suivi(config_test):
    """[revue W4] CAP-14 « Depuis l'accueil » : le choix du workflow et le formulaire incident
    sont dans `GET /`, le jeton CSRF est dans la page, et l'envoi des champs de ce formulaire
    (noms et action lus dans le HTML, jeton lu dans la page) aboutit au suivi de l'incident."""
    app = application(config_test)
    # Client sans jeton par défaut : le jeton envoyé est celui que porte la page.
    client = TestClient(app, base_url="http://testserver")
    accueil = client.get("/")
    assert accueil.status_code == 200
    verifier_accessibilite(accueil.text)
    arbre = lxml_html.fromstring(accueil.text)

    (choix,) = arbre.xpath('//*[@role="group"][@aria-label="Choix du workflow"]')
    cartes = {b.get("data-workflow"): b for b in choix.xpath(".//button")}
    assert set(cartes) == {"fiches", "incident"}
    assert cartes["fiches"].get("aria-pressed") == "true"
    assert cartes["incident"].get("aria-pressed") == "false"
    assert cartes["incident"].get("aria-controls") == "depot-incident"

    (formulaire,) = arbre.xpath('//form[@action="/incidents"]')
    assert formulaire.get("id") == "depot-incident"
    assert formulaire.get("method") == "post"
    assert formulaire.get("enctype") == "multipart/form-data"
    (champ_fichiers,) = formulaire.xpath('.//input[@type="file"]')
    (champ_modele,) = formulaire.xpath(".//select")
    (champ_annee,) = formulaire.xpath('.//input[@type="number"]')
    assert champ_fichiers.get("name") == "journaux"
    assert champ_fichiers.get("multiple") is not None
    assert (champ_annee.get("min"), champ_annee.get("max")) == ("1970", "2100")
    assert not formulaire.xpath(".//*[@name='langue']")  # pas de langue pour l'incident
    (option_defaut,) = champ_modele.xpath(".//option[@selected]")
    (jeton,) = arbre.xpath('//meta[@name="csrf"]/@content')
    assert jeton == app.state.jeton_csrf

    sans_jeton = client.post(
        formulaire.get("action"),
        files=[(champ_fichiers.get("name"), nom_contenu) for _, nom_contenu in _journaux_demo()],
        follow_redirects=False,
    )
    assert sans_jeton.status_code == 403
    assert sans_jeton.json()["code"] == "csrf_invalide"

    reponse = client.post(
        formulaire.get("action"),
        files=[(champ_fichiers.get("name"), nom_contenu) for _, nom_contenu in _journaux_demo()],
        data={
            champ_modele.get("name"): option_defaut.get("value"),
            champ_annee.get("name"): ANNEE,
        },
        headers={"X-CSRF-Token": jeton},
        follow_redirects=False,
    )
    assert reponse.status_code == 303
    id_lot = reponse.headers["location"].rsplit("/", 1)[-1]
    assert reponse.headers["location"] == f"/incidents/{id_lot}"
    suivi = client.get(reponse.headers["location"], follow_redirects=False)
    assert suivi.status_code == 200
    assert f'data-id-lot="{id_lot}"' in suivi.text
    assert _attendre_fin(client, id_lot)["etat"] == "termine"
    assert client.get(f"/incidents/{id_lot}/resultat").json()["total"] == 16


def test_accueil_lot_en_cours_desactive_aussi_le_depot_incident(config_test, faux_moteur):
    """[revue W4] EXPERIENCE « Accueil, lot déjà en cours » : dépôt désactivé, dans les deux
    workflows (verrou unique, §24)."""
    from tests.e2e.test_ui import _attendre_fin as _attendre_fin_fiches
    from tests.e2e.test_ui import _notes_demo

    client = client_web(_vers(config_test, faux_moteur("lent")))
    premier = client.post("/lots", files=_notes_demo()[:1], follow_redirects=False)
    id_lot = premier.headers["location"].rsplit("/", 1)[-1]
    arbre = lxml_html.fromstring(client.get("/").text)
    (formulaire,) = arbre.xpath('//form[@id="depot-incident"]')
    for element in formulaire.xpath(".//input | .//select | .//button"):
        assert element.get("disabled") is not None, element.get("name") or element.tag
    _attendre_fin_fiches(client, id_lot)
