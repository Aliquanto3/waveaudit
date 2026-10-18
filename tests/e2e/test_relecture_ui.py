"""API de relecture, régénération et export de bout en bout (contrats §24 à §27, W3a).

Faux moteur seulement ; minuteries courtes et attentes par sondage borné.
"""

import io
import json
import os
import re
import time
from datetime import datetime

import pytest
from docx import Document
from fastapi.testclient import TestClient
from lxml import html as lxml_html

from tests.accessibilite import verifier_accessibilite, verifier_focus_css
from tests.e2e.test_ui import DEMO, _attendre_fin
from tests.web_outils import client as client_web
from waveaudit.application import composition
from waveaudit.application import relecture as relecture_app
from waveaudit.application.config import racine_projet
from waveaudit.application.moteur_demande import Occupation
from waveaudit.domaine.erreurs import MoteurInjoignable

_RACINE_WEB = racine_projet() / "src" / "waveaudit" / "interfaces" / "web"

NOTE_SQL = "03-injection-sql.md"
NOTE_IDOR = "05-idor.md"
NOTE_SANS_RATTACHEMENT = (
    "zz-sans-rattachement.md",
    b"---\ntitre: Qwxz\n---\nqwxzv bnmkj plokiu ytreza wxcvbn\n",
)
NOTE_VIDE = ("zz-vide.md", b"---\ntitre: Vide\n---\n   \n")
VECTEUR_MOYEN = "CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:L/I:N/A:N"  # 5.3, moyenne


def _config(config, fm, **moteur):
    return config.model_copy(
        update={"moteur": config.moteur.model_copy(update={"base_url": fm.base_url, **moteur})}
    )


def _fichiers(*noms: str, extra: tuple[tuple[str, bytes], ...] = ()) -> list:
    demo = [("notes", (nom, (DEMO / nom).read_bytes())) for nom in noms]
    return demo + [("notes", (nom, contenu)) for nom, contenu in extra]


def _lancer(client: TestClient, fichiers: list, **donnees) -> str:
    reponse = client.post("/lots", files=fichiers, data=donnees, follow_redirects=False)
    assert reponse.status_code == 303, reponse.text
    return reponse.headers["location"].rsplit("/", 1)[-1]


def _lot_termine(client: TestClient, fichiers: list, **donnees) -> str:
    id_lot = _lancer(client, fichiers, **donnees)
    assert _attendre_fin(client, id_lot)["etat"] == "termine"
    return id_lot


def _entree(client: TestClient, id_lot: str, index: int) -> dict:
    reponse = client.get(f"/lots/{id_lot}/entrees/{index}")
    assert reponse.status_code == 200, reponse.text
    return reponse.json()


def _corps_correction(detail: dict, **champs) -> dict:
    fiche = detail["fiche_courante"]
    corps = {
        "version": detail["entree"]["version"],
        "titre": fiche["titre"],
        "description": fiche["description"],
        "impact": fiche["impact"],
        "recommandation": fiche["recommandation"],
        "cvss_vecteur": fiche["cvss"]["vecteur"],
        "references": [r["id"] for r in fiche["references"]],
    }
    return corps | champs


def _attendre_generation(client: TestClient, id_lot: str, index: int, timeout_s=10.0) -> dict:
    limite = time.monotonic() + timeout_s
    while True:
        detail = _entree(client, id_lot, index)
        if detail["entree"]["generation"] != "en_cours" and (
            client.app.state.verrou.occupation() is None
        ):
            return detail
        if time.monotonic() > limite:
            raise AssertionError(f"régénération de {id_lot}/{index} toujours en cours")
        time.sleep(0.05)


def _code(reponse, statut: int, code: str) -> dict:
    assert reponse.status_code == statut, reponse.text
    corps = reponse.json()
    assert corps["code"] == code
    assert ":\\" not in corps["detail"] and not corps["detail"].startswith("/")
    return corps


def _textes_docx(contenu: bytes) -> str:
    document = Document(io.BytesIO(contenu))
    lignes = [p.text for p in document.paragraphs]
    for table in document.tables:
        for rangee in table.rows:
            lignes.extend(cellule.text for cellule in rangee.cells)
    return "\n".join(lignes)


def _constat_demande(requete: dict) -> dict:
    user = [m for m in requete["messages"] if m["role"] == "user"][-1]["content"]
    debut = user.index("<constat>") + len("<constat>")
    return json.loads(user[debut : user.index("</constat>")])


# --- Lecture pendant le lot ------------------------------------------------------------------


def test_fiche_ouvrable_pendant_le_lot_progression_monotone(config_test, faux_moteur):
    lent = _config(config_test, faux_moteur("lent"), timeout_s=10.0)
    client = client_web(lent)
    id_lot = _lancer(client, _fichiers(NOTE_SQL, NOTE_IDOR))
    traitees: list[int] = []
    versions: list[list[int]] = []
    ouverte_pendant_le_lot = False
    limite = time.monotonic() + 20
    while time.monotonic() < limite:
        corps = client.get(f"/lots/{id_lot}/relecture").json()
        traitees.append(sum(1 for e in corps["entrees"] if e["etat"] != "attente"))
        versions.append([e["version"] for e in corps["entrees"]])
        if (
            corps["etat_lot"] == "en_cours"
            and corps["entrees"]
            and (corps["entrees"][0]["etat"] == "fiche")
        ):
            detail = _entree(client, id_lot, 1)
            assert detail["fiche_courante"]["titre"]
            assert detail["modifiable"] is True
            ouverte_pendant_le_lot = True
        if corps["etat_lot"] == "termine":
            break
        time.sleep(0.05)
    assert ouverte_pendant_le_lot
    assert traitees == sorted(traitees) and traitees[-1] == 2
    for avant, apres in zip(versions, versions[1:], strict=False):
        if len(avant) == len(apres):
            assert all(a <= b for a, b in zip(avant, apres, strict=True))


# --- Corrections et validation ---------------------------------------------------------------


def test_correction_sur_disque_relue_apres_recreation_cvss_du_serveur(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    detail = _entree(client, id_lot, 1)
    version = detail["entree"]["version"]
    corps = _corps_correction(detail, titre="Titre corrigé", cvss_vecteur=VECTEUR_MOYEN)

    reponse = client.put(f"/lots/{id_lot}/entrees/1/correction", json=corps)
    assert reponse.status_code == 200, reponse.text
    entree = reponse.json()["entree"]
    assert entree["version"] == version + 1
    assert entree["fiche_corrigee"]["cvss"] == {
        "vecteur": VECTEUR_MOYEN,
        "score": 5.3,
        "severite": "moyenne",
    }

    # Correction identique : ni écriture ni incrément (consigne R1/R2).
    identique = client.put(
        f"/lots/{id_lot}/entrees/1/correction", json=corps | {"version": version + 1}
    )
    assert identique.json()["entree"]["version"] == version + 1

    relu = _entree(client_web(config_test), id_lot, 1)
    assert relu["entree"]["version"] == version + 1
    assert relu["fiche_courante"]["titre"] == "Titre corrigé"
    assert relu["fiche_courante"]["cvss"]["score"] == 5.3


def test_vecteur_invalide_enregistre_puis_validation_refusee(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    detail = _entree(client, id_lot, 1)
    reponse = client.put(
        f"/lots/{id_lot}/entrees/1/correction",
        json=_corps_correction(detail, cvss_vecteur="CVSS:3.1/AV:X"),
    )
    assert reponse.status_code == 200
    entree = reponse.json()["entree"]
    assert entree["fiche_corrigee"]["cvss"] is None
    assert entree["fiche_corrigee"]["cvss_vecteur"] == "CVSS:3.1/AV:X"
    refus = client.post(f"/lots/{id_lot}/entrees/1/validation", json={"version": entree["version"]})
    _code(refus, 422, "vecteur_invalide")


def test_references_ajoutee_dupliquee_et_corps_invalide(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    detail = _entree(client, id_lot, 1)
    ids = [r["id"] for r in detail["fiche_courante"]["references"]]
    url = f"/lots/{id_lot}/entrees/1/correction"
    _code(
        client.put(url, json=_corps_correction(detail, references=["INVENTEE"])),
        422,
        "reference_ajoutee",
    )
    _code(
        client.put(url, json=_corps_correction(detail, references=ids[:1] * 2)),
        422,
        "reference_dupliquee",
    )
    _code(client.put(url, json={"version": 1}), 422, "corps_invalide")
    _code(
        client.put(url, json=_corps_correction(detail, cvss_vecteur="C" * 101)),
        422,
        "corps_invalide",
    )
    trop_long = "x" * (config_test.moteur.longueur_max_rubrique + 1)
    _code(
        client.put(url, json=_corps_correction(detail, impact=trop_long)),
        422,
        "rubrique_trop_longue",
    )
    assert _entree(client, id_lot, 1)["entree"]["version"] == detail["entree"]["version"]


def test_validee_puis_modifiee_revient_a_valider(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL, NOTE_IDOR))
    assert client.get(f"/lots/{id_lot}/relecture").json()["non_validees"] == 2
    detail = _entree(client, id_lot, 1)
    validee = client.post(
        f"/lots/{id_lot}/entrees/1/validation", json={"version": detail["entree"]["version"]}
    )
    assert validee.status_code == 200
    entree = validee.json()["entree"]
    assert entree["statut_relecture"] == "validee"
    assert client.get(f"/lots/{id_lot}/relecture").json()["non_validees"] == 1

    detail = _entree(client, id_lot, 1)
    modifiee = client.put(
        f"/lots/{id_lot}/entrees/1/correction",
        json=_corps_correction(detail, titre="Autre titre"),
    )
    assert modifiee.json()["entree"]["statut_relecture"] == "a_valider"


def test_conflit_version_renvoie_l_entree_a_jour(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    detail = _entree(client, id_lot, 1)
    url = f"/lots/{id_lot}/entrees/1/correction"
    assert client.put(url, json=_corps_correction(detail, titre="Premier")).status_code == 200
    conflit = client.put(url, json=_corps_correction(detail, titre="Second"))
    corps = _code(conflit, 409, "conflit_version")
    assert corps["entree"]["version"] == detail["entree"]["version"] + 1
    assert corps["entree"]["fiche_corrigee"]["titre"] == "Premier"
    _code(
        client.post(f"/lots/{id_lot}/entrees/1/validation", json={"version": 1}),
        409,
        "conflit_version",
    )


# --- Changement de candidat ------------------------------------------------------------------


def test_changement_de_candidat_confirmation_et_garde_fous(config_test, faux_moteur):
    fm = faux_moteur()
    config = _config(config_test, fm)
    client = client_web(config)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    detail = _entree(client, id_lot, 1)
    assert detail["generable"] is True
    assert len(detail["candidats_detail"]) == 3
    assert all(c["titre"] for c in detail["candidats_detail"])
    courant = detail["entree"]["candidat_choisi"]
    autre = next(c["id"] for c in detail["entree"]["candidats"] if c["id"] != courant)
    url = f"/lots/{id_lot}/entrees/1/candidat"

    _code(
        client.post(url, json={"version": detail["entree"]["version"], "constat_id": courant}),
        422,
        "candidat_courant",
    )
    _code(
        client.post(url, json={"version": detail["entree"]["version"], "constat_id": "XX-NON"}),
        422,
        "candidat_inconnu",
    )
    validee = client.post(
        f"/lots/{id_lot}/entrees/1/validation", json={"version": detail["entree"]["version"]}
    ).json()["entree"]
    _code(
        client.post(url, json={"version": validee["version"], "constat_id": autre}),
        409,
        "confirmation_requise",
    )
    assert client.app.state.verrou.occupation() is None  # tout refus libère le verrou

    requetes_avant = len(fm.requetes)
    lance = client.post(
        url, json={"version": validee["version"], "constat_id": autre, "confirmer": True}
    )
    assert lance.status_code == 202, lance.text
    assert lance.json()["entree"]["generation"] == "en_cours"
    apres = _attendre_generation(client, id_lot, 1)
    entree = apres["entree"]
    assert entree["generation"] == "aucune"
    assert entree["candidat_choisi"] == autre
    assert entree["statut_relecture"] == "a_valider"
    assert entree["fiche_corrigee"] is None
    assert entree["version"] == validee["version"] + 2
    assert entree["fiche_generee"]["tracabilite"]["constat_id"] == autre
    assert entree["fiche_generee"]["statut"] in ("ok", "a_revoir")
    assert len(fm.requetes) > requetes_avant
    assert _constat_demande(fm.requetes[-1])["id"] == autre


def test_candidat_et_export_refuses_pendant_un_lot(config_test, faux_moteur):
    lent = _config(config_test, faux_moteur("lent"), timeout_s=10.0)
    client = client_web(lent)
    premier = _lot_termine(client, _fichiers(NOTE_SQL))
    detail = _entree(client, premier, 1)
    autre = next(
        c["id"]
        for c in detail["entree"]["candidats"]
        if c["id"] != detail["entree"]["candidat_choisi"]
    )
    second = _lancer(client, _fichiers(NOTE_IDOR))
    refus = client.post(
        f"/lots/{premier}/entrees/1/candidat",
        json={"version": detail["entree"]["version"], "constat_id": autre},
    )
    corps = _code(refus, 409, "lot_en_cours")
    assert corps["occupation"]["id_lot"] == second
    _code(client.post(f"/lots/{second}/export", json={"type": "brouillon"}), 409, "lot_en_cours")
    assert _entree(client, premier, 1)["entree"] == detail["entree"]
    _attendre_fin(client, second)


def test_non_rattachee_vers_fiche_et_non_traitee_non_generable(config_test, faux_moteur):
    fm = faux_moteur()
    client = client_web(_config(config_test, fm))
    id_lot = _lot_termine(client, _fichiers(extra=(NOTE_SANS_RATTACHEMENT, NOTE_VIDE)))
    resume = client.get(f"/lots/{id_lot}/relecture").json()["entrees"]
    assert [(e["etat"], e["statut_generation"]) for e in resume] == [
        ("ecartee", "non_rattachee"),
        ("ecartee", "non_traitee"),
    ]
    vide = _entree(client, id_lot, 2)
    assert vide["generable"] is False
    _code(
        client.post(
            f"/lots/{id_lot}/entrees/2/candidat",
            json={"version": vide["entree"]["version"], "constat_id": "WEB-INJECTION-SQL"},
        ),
        422,
        "non_generable",
    )

    detail = _entree(client, id_lot, 1)
    assert detail["generable"] is True
    candidats = {c["id"] for c in detail["entree"]["candidats"]}
    catalogue = client.get(f"/lots/{id_lot}/catalogue").json()["constats"]
    choisi = next(c["id"] for c in catalogue if c["id"] not in candidats)
    url = f"/lots/{id_lot}/entrees/1/candidat"
    _code(
        client.post(url, json={"version": detail["entree"]["version"], "constat_id": "XX-NON"}),
        422,
        "candidat_inconnu",
    )
    lance = client.post(url, json={"version": detail["entree"]["version"], "constat_id": choisi})
    assert lance.status_code == 202, lance.text
    entree = _attendre_generation(client, id_lot, 1)["entree"]
    assert entree["etat"] == "fiche"
    assert entree["candidat_choisi"] == choisi
    assert entree["ecartee"] is None
    assert _constat_demande(fm.requetes[-1])["id"] == choisi


# --- Export ----------------------------------------------------------------------------------


def test_export_final_et_brouillon_relus_par_python_docx(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL, NOTE_IDOR))
    url = f"/lots/{id_lot}/export"
    _code(client.post(url, json={"type": "final"}), 409, "aucune_fiche_validee")
    _code(client.post(url, json={"type": "autre"}), 422, "type_invalide")
    _code(client.post(url, content=b"pas du json"), 422, "type_invalide")

    detail = _entree(client, id_lot, 1)
    corrigee = client.put(
        f"/lots/{id_lot}/entrees/1/correction",
        json=_corps_correction(detail, titre="Injection SQL corrigée"),
    ).json()["entree"]
    client.post(f"/lots/{id_lot}/entrees/1/validation", json={"version": corrigee["version"]})
    titre_non_valide = _entree(client, id_lot, 2)["fiche_courante"]["titre"]

    brouillon = client.post(url, json={"type": "brouillon"})
    assert brouillon.status_code == 200, brouillon.text
    assert f"brouillon-{id_lot}.docx" in brouillon.headers["content-disposition"]
    assert "x-non-validees" not in brouillon.headers
    texte = _textes_docx(brouillon.content)
    assert "Injection SQL corrigée" in texte
    assert f"{titre_non_valide} — Non validée" in texte

    final = client.post(url, json={"type": "final"})
    assert final.status_code == 200, final.text
    assert f"rapport-{id_lot}.docx" in final.headers["content-disposition"]
    assert final.headers["x-non-validees"] == "1"
    texte = _textes_docx(final.content)
    assert "Injection SQL corrigée" in texte
    assert titre_non_valide not in texte
    assert NOTE_IDOR in texte and "Non validée" in texte  # annexe des non validées
    assert (config_test.chemins.var / "sorties" / id_lot / f"rapport-{id_lot}.docx").is_file()


def test_get_export_405_sans_fichier_ecrit(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    dossier = config_test.chemins.var / "sorties" / id_lot
    avant = sorted(p.name for p in dossier.iterdir())
    _code(client.get(f"/lots/{id_lot}/export"), 405, "methode_refusee")
    assert sorted(p.name for p in dossier.iterdir()) == avant


def test_export_verrouille_409(config_test, monkeypatch):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))

    class _Verrouille:
        def exporter(self, *args, **kwargs):
            raise PermissionError("fichier ouvert")

    monkeypatch.setattr(composition, "creer_exporteur", lambda config: _Verrouille())
    _code(
        client.post(f"/lots/{id_lot}/export", json={"type": "brouillon"}), 409, "export_verrouille"
    )


def test_lot_echoue_relisible_et_exportable(config_test, monkeypatch):
    from waveaudit.workflows.fiches import pipeline

    traiter = pipeline.WorkflowFiches.traiter
    appels: list[str] = []

    def traiter_coupe(self, element):
        appels.append(element.fichier)
        if len(appels) == 2:
            raise MoteurInjoignable("moteur coupé")
        return traiter(self, element)

    monkeypatch.setattr(pipeline.WorkflowFiches, "traiter", traiter_coupe)
    client = client_web(config_test)
    id_lot = _lancer(client, _fichiers(NOTE_SQL, NOTE_IDOR))
    assert _attendre_fin(client, id_lot)["etat"] == "echoue"
    corps = client.get(f"/lots/{id_lot}/relecture").json()
    assert corps["etat_lot"] == "echoue"
    assert [e["etat"] for e in corps["entrees"]] == ["fiche", "attente"]
    assert _entree(client, id_lot, 2)["modifiable"] is False
    brouillon = client.post(f"/lots/{id_lot}/export", json={"type": "brouillon"})
    assert brouillon.status_code == 200
    assert NOTE_IDOR in _textes_docx(brouillon.content)  # annexe : lot interrompu


# --- Revue : catalogue, orphelines, autre processus, écriture, bornes, modèle ---------------


def test_catalogue_trie_et_dans_la_langue_du_lot(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL), langue="en")
    constats = client.get(f"/lots/{id_lot}/catalogue").json()["constats"]
    attendus = composition.creer_catalogue(config_test).constats()
    assert [c["id"] for c in constats] == sorted(c.id for c in attendus)
    par_id = {c.id: c for c in attendus}
    assert all(c["titre"] == par_id[c["id"]].en.titre for c in constats)
    assert all(c["famille"] == par_id[c["id"]].famille for c in constats)
    detail = _entree(client, id_lot, 1)
    assert all(c["titre"] == par_id[c["id"]].en.titre for c in detail["candidats_detail"])
    _code(client.get("/lots/inconnu/catalogue"), 404, "lot_inconnu")


def test_generation_en_cours_orpheline_lue_sans_ecriture(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    chemin = config_test.chemins.var / "sorties" / id_lot / "relecture.json"
    donnees = json.loads(chemin.read_text(encoding="utf-8"))
    donnees["entrees"][0]["generation"] = "en_cours"
    chemin.write_text(json.dumps(donnees), encoding="utf-8")
    brut = chemin.read_bytes()
    version = donnees["entrees"][0]["version"]

    detail = _entree(client, id_lot, 1)
    assert detail["entree"]["generation"] == "echec"
    assert detail["entree"]["erreur_generation"] == "génération interrompue"
    assert detail["entree"]["version"] == version
    resume = client.get(f"/lots/{id_lot}/relecture").json()["entrees"][0]
    assert resume["generation"] == "echec"
    assert chemin.read_bytes() == brut
    # Relançable : l'entrée interrompue accepte un nouveau candidat.
    autre = next(
        c["id"]
        for c in detail["entree"]["candidats"]
        if c["id"] != detail["entree"]["candidat_choisi"]
    )
    lance = client.post(
        f"/lots/{id_lot}/entrees/1/candidat", json={"version": version, "constat_id": autre}
    )
    assert lance.status_code == 202, lance.text
    assert _attendre_generation(client, id_lot, 1)["entree"]["candidat_choisi"] == autre


def test_lot_en_cours_dans_un_autre_processus_409(config_test, monkeypatch):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    dossier = config_test.chemins.var / "sorties" / id_lot
    fiche_lot = json.loads((dossier / "lot.json").read_text(encoding="utf-8"))
    fiche_lot["pid"] = os.getpid() + 1
    (dossier / "lot.json").write_text(json.dumps(fiche_lot), encoding="utf-8")
    lignes = (dossier / "journal.jsonl").read_text(encoding="utf-8").splitlines()
    sans_fin = [ligne for ligne in lignes if json.loads(ligne)["type"] != "lot_termine"]
    (dossier / "journal.jsonl").write_text("\n".join(sans_fin) + "\n", encoding="utf-8")
    monkeypatch.setattr(composition, "processus_actif", lambda pid, cree_le: True)

    detail = _entree(client, id_lot, 1)
    assert detail["modifiable"] is False and detail["generable"] is False
    assert client.get(f"/lots/{id_lot}/relecture").json()["etat_lot"] == "en_cours"
    version = detail["entree"]["version"]
    base = f"/lots/{id_lot}/entrees/1"
    _code(
        client.put(f"{base}/correction", json=_corps_correction(detail, titre="x")),
        409,
        "lot_en_cours_externe",
    )
    _code(client.post(f"{base}/validation", json={"version": version}), 409, "lot_en_cours_externe")
    autre = detail["entree"]["candidats"][1]["id"]
    _code(
        client.post(f"{base}/candidat", json={"version": version, "constat_id": autre}),
        409,
        "lot_en_cours_externe",
    )
    _code(
        client.post(f"/lots/{id_lot}/export", json={"type": "brouillon"}),
        409,
        "lot_en_cours_externe",
    )
    assert _entree(client, id_lot, 1)["entree"]["version"] == version


def test_os_replace_en_echec_503_sans_changement(config_test, monkeypatch):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    chemin = config_test.chemins.var / "sorties" / id_lot / "relecture.json"
    brut = chemin.read_bytes()
    detail = _entree(client, id_lot, 1)

    def replace_refuse(source, cible):
        raise PermissionError("fichier tenu par un autre programme")

    monkeypatch.setattr(relecture_app, "PAUSE_REMPLACEMENT_S", 0.0)
    monkeypatch.setattr(relecture_app.os, "replace", replace_refuse)
    _code(
        client.put(
            f"/lots/{id_lot}/entrees/1/correction",
            json=_corps_correction(detail, titre="Jamais écrit"),
        ),
        503,
        "ecriture_impossible",
    )
    _code(
        client.post(
            f"/lots/{id_lot}/entrees/1/validation", json={"version": detail["entree"]["version"]}
        ),
        503,
        "ecriture_impossible",
    )
    assert chemin.read_bytes() == brut


@pytest.mark.parametrize("index", [0, 2])
def test_index_hors_bornes_404(config_test, index):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    _code(client.get(f"/lots/{id_lot}/entrees/{index}"), 404, "entree_inconnue")
    _code(
        client.post(f"/lots/{id_lot}/entrees/{index}/validation", json={"version": 1}),
        404,
        "entree_inconnue",
    )


def test_modele_desinstalle_422_modele_indisponible(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    local = config_test.model_copy(
        update={"moteur": config_test.moteur.model_copy(update={"mode": "local-gere"})}
    )
    assert composition.modeles_installes(local) == []
    client_local = client_web(local)
    detail = _entree(client_local, id_lot, 1)
    assert detail["generable"] is False
    autre = detail["entree"]["candidats"][1]["id"]
    chemin = config_test.chemins.var / "sorties" / id_lot / "relecture.json"
    brut = chemin.read_bytes()
    _code(
        client_local.post(
            f"/lots/{id_lot}/entrees/1/candidat",
            json={"version": detail["entree"]["version"], "constat_id": autre},
        ),
        422,
        "modele_indisponible",
    )
    assert chemin.read_bytes() == brut
    assert client_local.app.state.verrou.occupation() is None


def test_relecture_illisible_erreur_codee(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    chemin = config_test.chemins.var / "sorties" / id_lot / "relecture.json"
    chemin.write_text("{pas du json", encoding="utf-8")
    _code(client.get(f"/lots/{id_lot}/relecture"), 409, "relecture_illisible")
    _code(client.get(f"/lots/{id_lot}/entrees/1"), 409, "relecture_illisible")
    _code(
        client.post(f"/lots/{id_lot}/entrees/1/validation", json={"version": 1}),
        409,
        "relecture_illisible",
    )
    _code(
        client.post(f"/lots/{id_lot}/export", json={"type": "brouillon"}),
        409,
        "relecture_illisible",
    )


def test_lot_inconnu_et_lot_incident_404(config_test):
    client = client_web(config_test)
    _code(client.get("/lots/inconnu/relecture"), 404, "lot_inconnu")
    _code(client.post("/lots/inconnu/export", json={"type": "final"}), 404, "lot_inconnu")
    dossier = config_test.chemins.var / "sorties" / "incident-x"
    dossier.mkdir(parents=True)
    (dossier / "lot.json").write_text(
        json.dumps(
            {
                "id_lot": "incident-x",
                "workflow": "incident",
                "langue": None,
                "modele_demande": None,
                "cree_le": "2020-01-01T00:00:00+00:00",
                "pid": 999999999,
            }
        ),
        encoding="utf-8",
    )
    _code(client.get("/lots/incident-x/relecture"), 404, "lot_inconnu")
    _code(client.post("/lots/incident-x/export", json={"type": "final"}), 404, "lot_incident")


def test_regeneration_occupee_par_ce_processus_409(config_test):
    """Le verrou d'une régénération en cours refuse un second candidat (ordre du §26)."""
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    detail = _entree(client, id_lot, 1)
    verrou = client.app.state.verrou
    assert verrou.prendre(Occupation("regeneration", id_lot, 1, datetime.now().astimezone()))
    try:
        refus = client.post(
            f"/lots/{id_lot}/entrees/1/candidat",
            json={
                "version": detail["entree"]["version"],
                "constat_id": detail["entree"]["candidats"][1]["id"],
            },
        )
        corps = _code(refus, 409, "lot_en_cours")
        assert corps["occupation"]["nature"] == "regeneration"
        _code(
            client.post(f"/lots/{id_lot}/export", json={"type": "brouillon"}), 409, "lot_en_cours"
        )
    finally:
        verrou.liberer()


# --- Revue W3a -------------------------------------------------------------------------------


def test_json_trop_imbrique_refus_code(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    imbrique = "[" * 50_000
    entetes = {"Content-Type": "application/json"}
    _code(
        client.post(f"/lots/{id_lot}/entrees/1/validation", content=imbrique, headers=entetes),
        422,
        "corps_invalide",
    )
    _code(
        client.post(f"/lots/{id_lot}/export", content=imbrique, headers=entetes),
        422,
        "type_invalide",
    )


def test_relecture_verrouillee_en_lecture_409(config_test, monkeypatch):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    lire = type(config_test.chemins.var).read_text

    def lire_verrouille(chemin, *args, **kwargs):
        if chemin.name == "relecture.json":
            raise PermissionError("verrouillé")
        return lire(chemin, *args, **kwargs)

    monkeypatch.setattr(type(config_test.chemins.var), "read_text", lire_verrouille)
    _code(client.get(f"/lots/{id_lot}/relecture"), 409, "relecture_illisible")
    _code(
        client.post(f"/lots/{id_lot}/entrees/1/candidat", json={"version": 1, "constat_id": "x"}),
        409,
        "relecture_illisible",
    )
    assert client.app.state.verrou.occupation() is None


def test_docx_verrouille_a_la_relecture_409(config_test, monkeypatch):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    lire = type(config_test.chemins.var).read_bytes

    def lire_verrouille(chemin):
        if chemin.name.startswith("brouillon-"):
            raise PermissionError("ouvert dans Word")
        return lire(chemin)

    monkeypatch.setattr(type(config_test.chemins.var), "read_bytes", lire_verrouille)
    _code(
        client.post(f"/lots/{id_lot}/export", json={"type": "brouillon"}), 409, "export_verrouille"
    )


def test_lot_sans_modele_consigne_regenerable_sur_le_modele_par_defaut(config_test, monkeypatch):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    chemin = config_test.chemins.var / "sorties" / id_lot / "relecture.json"
    etat = json.loads(chemin.read_text(encoding="utf-8"))
    chemin.write_text(json.dumps(etat | {"modele": ""}), encoding="utf-8")
    local = config_test.model_copy(
        update={"moteur": config_test.moteur.model_copy(update={"mode": "local-gere"})}
    )
    installe = composition.ModeleInstalle("gemma", "modele_defaut", 1, None, True)
    monkeypatch.setattr(composition, "modeles_installes", lambda config: [installe])
    assert _entree(client_web(local), id_lot, 1)["generable"] is True
    monkeypatch.setattr(composition, "modeles_installes", lambda config: [])
    assert _entree(client_web(local), id_lot, 1)["generable"] is False


# --- Écran de relecture (W3b, contrats §26, §28 [revue V2]) -----------------------------------


def test_accessibilite_relecture(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    page = client.get(f"/lots/{id_lot}")
    assert page.status_code == 200
    verifier_accessibilite(page.text)
    verifier_focus_css((_RACINE_WEB / "static" / "style.css").read_text(encoding="utf-8"))


def test_tiroir_listbox_et_barre_candidats_tablist(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    arbre = lxml_html.fromstring(client.get(f"/lots/{id_lot}").text)

    (tiroir,) = arbre.xpath('//*[@role="listbox"][@aria-label="Fiches du lot"]')
    assert tiroir.xpath('.//*[@role="option"]')

    (tablist,) = arbre.xpath('//*[@role="tablist"]')
    onglets = tablist.xpath('.//*[@role="tab"]')
    assert len(onglets) == 3
    for onglet in onglets:
        assert (onglet.get("aria-label") or onglet.text_content()).strip()
        assert onglet.get("aria-controls")

    (panneau,) = arbre.xpath('//*[@role="tabpanel"]')
    assert panneau.get("id") in tablist.xpath('.//*[@role="tab"]/@aria-controls')


def test_liste_du_catalogue_avec_libelle_et_filtre(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    arbre = lxml_html.fromstring(client.get(f"/lots/{id_lot}").text)

    (filtre,) = arbre.xpath('//input[@id="catalogue-filtre"]')
    (label,) = arbre.xpath('//label[@for="catalogue-filtre"]')
    assert label.text_content().strip()
    (catalogue,) = arbre.xpath('//*[@role="listbox"][@aria-label="Constats du catalogue"]')
    assert catalogue is not None


def test_dialogues_de_confirmation_et_export_final(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    arbre = lxml_html.fromstring(client.get(f"/lots/{id_lot}").text)

    dialogues = arbre.xpath('//*[@role="dialog" or @role="alertdialog"]')
    assert len(dialogues) >= 3
    for dialogue in dialogues:
        assert dialogue.get("aria-modal") == "true"
        (titre,) = arbre.xpath("//*[@id=$i]", i=dialogue.get("aria-labelledby"))
        assert titre.text_content().strip()

    (export,) = arbre.xpath('//*[@id="dialog-export"]')
    assert export.get("role") in ("dialog", "alertdialog")
    assert export.get("aria-modal") == "true"


def test_lang_sur_panneau_note_et_formulaire(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    arbre = lxml_html.fromstring(client.get(f"/lots/{id_lot}").text)

    (panneau_note,) = arbre.xpath('//*[@id="panneau-note"]')
    (panneau_fiche,) = arbre.xpath('//*[@id="panneau-fiche"]')
    assert (panneau_note.get("lang") or "").strip()
    assert (panneau_fiche.get("lang") or "").strip()
    assert panneau_fiche.xpath('.//form[@id="formulaire-fiche"]')


def test_aucun_attribut_on_ou_style_en_ligne_sur_la_page_de_relecture(config_test):
    client = client_web(config_test)
    id_lot = _lot_termine(client, _fichiers(NOTE_SQL))
    texte = client.get(f"/lots/{id_lot}").text
    assert re.search(r"\son\w+\s*=", texte, re.IGNORECASE) is None
    assert re.search(r"\sstyle\s*=", texte, re.IGNORECASE) is None
    assert "<style" not in texte.lower()
