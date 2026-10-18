"""Tests de l'export Word (CAP-6, contrats §17, AD-14)."""

import inspect
import re
import typing
from pathlib import Path

import pytest
from docx import Document

from waveaudit.adaptateurs.export_docx import (
    BANDEAU,
    BANDEAUX,
    DETAILS_EN,
    LIBELLES,
    MOTIFS_EN,
    RUBRIQUES_EN,
    STATUTS_ECARTES_EN,
    ExportDocx,
    _motif_affiche,
)
from waveaudit.domaine.erreurs import ConfigurationInvalide
from waveaudit.domaine.modeles import (
    RUBRIQUES,
    Candidat,
    CodeMotif,
    Cvss,
    EntreeEcartee,
    Fiche,
    Motif,
    Rapport,
    Reference,
    Statut,
    Tracabilite,
)
from waveaudit.domaine.ports import OptionsExport
from waveaudit.domaine.qualite_texte import motif_langue
from waveaudit.workflows.fiches import pipeline

_REFERENCE = Reference(
    id="CWE-79", titre="XSS", url="https://cwe.mitre.org/data/definitions/79.html"
)


def _fiche_ok() -> Fiche:
    return Fiche(
        titre="Injection SQL sur le formulaire de connexion",
        description="La requête concatène l'entrée utilisateur sans échappement.",
        impact="Contournement de l'authentification, lecture de la base.",
        recommandation="Utiliser des requêtes paramétrées.",
        cvss=Cvss(
            vecteur="CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H", score=9.8, severite="critique"
        ),
        references=[_REFERENCE],
        statut=Statut.OK,
        motifs=[],
        tracabilite=Tracabilite(
            note="01-injection-sql.md",
            constat_id="WEB-INJECTION-SQL",
            candidats=[Candidat(id="WEB-INJECTION-SQL", score=0.9)],
            constat_attendu="WEB-INJECTION-SQL",
        ),
    )


def _fiche_a_revoir() -> Fiche:
    return Fiche(
        titre="Vecteur CVSS invalide",
        description="d",
        impact="i",
        recommandation="r",
        cvss=Cvss(
            vecteur="CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:N/I:N/A:N", score=0.0, severite="aucune"
        ),
        references=[_REFERENCE],
        statut=Statut.A_REVOIR,
        motifs=[
            Motif(
                code=CodeMotif.VECTEUR_INVALIDE,
                detail="CVSS:3.1/AV:N/AC:X/PR:N/UI:N/S:U/C:N/I:N/A:N",
            )
        ],
        tracabilite=Tracabilite(
            note="02-notes.md",
            constat_id="WEB-XSS-REFLECHI",
            candidats=[Candidat(id="WEB-XSS-REFLECHI", score=0.4)],
            constat_attendu=None,
        ),
    )


def _rapport(avec_ecartees: bool) -> Rapport:
    ecartees = (
        [
            EntreeEcartee(
                note="03-hors-sujet.md",
                statut=Statut.NON_RATTACHEE,
                motif=Motif(
                    code=CodeMotif.AUCUN_CONSTAT, detail="aucun candidat au-dessus du seuil"
                ),
                candidats=[],
                constat_attendu=None,
            ),
            EntreeEcartee(
                note="04-invalide.md",
                statut=Statut.NON_TRAITEE,
                motif=Motif(code=CodeMotif.NOTE_INVALIDE, detail="corps vide"),
                candidats=[],
                constat_attendu=None,
            ),
        ]
        if avec_ecartees
        else []
    )
    return Rapport(fiches=[_fiche_ok(), _fiche_a_revoir()], ecartees=ecartees)


def _textes(document: Document) -> list[str]:
    return [p.text for p in document.paragraphs]


def test_theme_generique_sans_dossier_produit_un_rapport_complet(tmp_path: Path):
    chemin = tmp_path / "sortie" / "rapport.docx"
    ExportDocx(tmp_path / "theme-prive-absent").exporter(_rapport(avec_ecartees=True), chemin)

    assert chemin.is_file()
    document = Document(str(chemin))

    for section in document.sections:
        assert section.header.paragraphs[0].text == BANDEAU
        assert section.footer.paragraphs[0].text == BANDEAU

    titres = [
        p.text
        for p in document.paragraphs
        if p.style.name.startswith("Heading") or p.style.name == "Title"
    ]
    assert "Rapport de constats" in titres
    assert "1. Injection SQL sur le formulaire de connexion" in titres
    assert "2. Vecteur CVSS invalide" in titres
    intitules = (
        "Description",
        "Impact",
        "Recommandation",
        "CVSS 3.1",
        "Références",
        "Statut",
        "Traçabilité",
    )
    for intitule in intitules:
        assert titres.count(intitule) == 2
    assert "Annexe — notes écartées" in titres

    textes = _textes(document)
    assert "À revoir" in textes
    assert "vecteur_invalide : CVSS:3.1/AV:N/AC:X/PR:N/UI:N/S:U/C:N/I:N/A:N" in textes

    tables = document.tables
    assert len(tables) == 2
    synthese = [cellule.text for cellule in tables[0].rows[0].cells]
    assert synthese == ["N°", "Titre", "Sévérité", "Score", "Statut"]
    ligne_ok = [cellule.text for cellule in tables[0].rows[1].cells]
    assert ligne_ok[-1] == "OK"
    ligne_a_revoir = [cellule.text for cellule in tables[0].rows[2].cells]
    assert ligne_a_revoir[-1] == "À revoir"

    annexe = [cellule.text for cellule in tables[1].rows[0].cells]
    assert annexe == ["Note", "Statut", "Motif"]
    ligne_1 = [c.text for c in tables[1].rows[1].cells]
    assert ligne_1 == [
        "03-hors-sujet.md",
        "non_rattachee",
        "aucun_constat : aucun candidat au-dessus du seuil",
    ]
    ligne_2 = [c.text for c in tables[1].rows[2].cells]
    assert ligne_2 == ["04-invalide.md", "non_traitee", "note_invalide : corps vide"]


def test_sans_ecartees_pas_dannexe(tmp_path: Path):
    chemin = tmp_path / "rapport.docx"
    ExportDocx(tmp_path / "theme-prive-absent").exporter(_rapport(avec_ecartees=False), chemin)

    document = Document(str(chemin))
    titres = [p.text for p in document.paragraphs]
    assert "Annexe — notes écartées" not in titres
    assert len(document.tables) == 1


def test_theme_prive_avec_modele_et_theme_toml(tmp_path: Path):
    theme_prive = tmp_path / "theme-prive"
    theme_prive.mkdir()
    Document().save(str(theme_prive / "modele.docx"))
    (theme_prive / "theme.toml").write_text(
        'couleur_primaire = "#123456"\ncle_inconnue = "ignoree"\n', encoding="utf-8"
    )

    chemin = tmp_path / "rapport.docx"
    ExportDocx(theme_prive).exporter(_rapport(avec_ecartees=False), chemin)

    document = Document(str(chemin))
    assert document.styles["Heading 1"].font.color.rgb is not None
    assert str(document.styles["Heading 1"].font.color.rgb) == "123456"
    titres = [p.text for p in document.paragraphs]
    assert "Rapport de constats" in titres


def _modele_sans_styles_de_titre(chemin: Path, entete: str) -> None:
    """Modèle comme Word en écrit : seuls les styles employés sont définis (ni Title ni Heading)."""
    document = Document()
    attribut = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}styleId"
    for style in list(document.styles.element):
        identifiant = style.get(attribut) or ""
        if identifiant.startswith("Heading") or identifiant in ("Title", "TableGrid"):
            document.styles.element.remove(style)
    document.styles["Normal"].font.name = "Police Du Modele"
    document.sections[0].different_first_page_header_footer = True
    document.sections[0].header.paragraphs[0].text = entete
    document.save(str(chemin))


def test_theme_prive_via_config_modele_sans_styles_de_titre(tmp_path: Path, config_test):
    theme_prive = config_test.chemins.theme_prive
    theme_prive.mkdir(parents=True)
    _modele_sans_styles_de_titre(theme_prive / "modele.docx", "Mention du modèle")

    chemin = tmp_path / "rapport.docx"
    ExportDocx(config_test.chemins.theme_prive).exporter(_rapport(avec_ecartees=True), chemin)

    document = Document(str(chemin))
    assert document.styles["Normal"].font.name == "Police Du Modele"
    section = document.sections[0]
    assert [p.text for p in section.header.paragraphs] == ["Mention du modèle", BANDEAU]
    for conteneur in (section.footer, section.first_page_header, section.first_page_footer):
        assert BANDEAU in [p.text for p in conteneur.paragraphs]
    titres = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert titres.count("Description") == 2
    assert "Annexe — notes écartées" in titres


def test_caracteres_interdits_en_xml_retires(tmp_path: Path):
    fiche = _fiche_ok().model_copy(update={"description": "avant\x00\x0bapres", "titre": "T\x1f"})
    chemin = tmp_path / "rapport.docx"
    ExportDocx(tmp_path / "absent").exporter(Rapport(fiches=[fiche], ecartees=[]), chemin)

    textes = _textes(Document(str(chemin)))
    assert "avantapres" in textes
    assert "1. T" in textes
    assert [p.name for p in tmp_path.iterdir()] == ["rapport.docx"]


@pytest.mark.parametrize(
    "contenu",
    ['couleur_primaire = "rouge"', 'couleur_secondaire = "#12345"', "couleur_primaire = 3", "x ="],
)
def test_theme_toml_invalide_leve_configuration_invalide(tmp_path: Path, contenu: str):
    theme_prive = tmp_path / "theme-prive"
    theme_prive.mkdir()
    (theme_prive / "theme.toml").write_text(contenu, encoding="utf-8")

    with pytest.raises(ConfigurationInvalide, match="theme.toml"):
        ExportDocx(theme_prive)


def test_modele_docx_illisible_leve_configuration_invalide(tmp_path: Path):
    theme_prive = tmp_path / "theme-prive"
    theme_prive.mkdir()
    (theme_prive / "modele.docx").write_bytes(b"pas un docx")

    with pytest.raises(ConfigurationInvalide, match="modele.docx"):
        ExportDocx(theme_prive)


def test_theme_toml_sans_modele_colore_le_theme_generique(tmp_path: Path):
    theme_prive = tmp_path / "theme-prive"
    theme_prive.mkdir()
    (theme_prive / "theme.toml").write_text('couleur_secondaire = "#abcdef"', encoding="utf-8")

    chemin = tmp_path / "rapport.docx"
    ExportDocx(theme_prive).exporter(_rapport(avec_ecartees=False), chemin)

    document = Document(str(chemin))
    assert str(document.styles["Heading 3"].font.color.rgb) == "ABCDEF"
    assert str(document.styles["Heading 1"].font.color.rgb) == "1F3864"


def _titres(document: Document) -> list[str]:
    return [
        p.text
        for p in document.paragraphs
        if p.style.name.startswith("Heading") or p.style.name == "Title"
    ]


def test_en_bandeau_intitules_severites_statuts_annexe(tmp_path: Path):
    chemin = tmp_path / "rapport.docx"
    ExportDocx(tmp_path / "absent").exporter(
        _rapport(avec_ecartees=True), chemin, OptionsExport(langue="en")
    )

    document = Document(str(chemin))
    libelles = LIBELLES["en"]
    for section in document.sections:
        assert section.header.paragraphs[0].text == BANDEAUX["en"]
        assert section.footer.paragraphs[0].text == BANDEAUX["en"]

    titres = _titres(document)
    assert libelles["titre"] in titres
    intitules = (
        "description",
        "impact",
        "recommandation",
        "cvss",
        "references",
        "statut",
        "tracabilite",
    )
    for intitule in intitules:
        assert titres.count(libelles[intitule]) == 2
    assert libelles["annexe_titre"] in titres

    textes = _textes(document)
    assert "To review" in textes
    assert any(
        "Invalid CVSS vector : CVSS:3.1/AV:N/AC:X/PR:N/UI:N/S:U/C:N/I:N/A:N" in t for t in textes
    )
    assert not any("vecteur_invalide" in t for t in textes)

    tables = document.tables
    synthese_entetes = [cellule.text for cellule in tables[0].rows[0].cells]
    assert synthese_entetes == list(libelles["synthese"])
    ligne_critique = [cellule.text for cellule in tables[0].rows[1].cells]
    assert ligne_critique[2] == "Critical"
    ligne_a_revoir = [cellule.text for cellule in tables[0].rows[2].cells]
    assert ligne_a_revoir[2] == "None"
    assert ligne_a_revoir[-1] == "To review"

    annexe_entetes = [cellule.text for cellule in tables[1].rows[0].cells]
    assert annexe_entetes == list(libelles["annexe_entetes"])
    ligne_1 = [c.text for c in tables[1].rows[1].cells]
    assert ligne_1 == ["03-hors-sujet.md", "Not matched", "No matching finding"]
    ligne_2 = [c.text for c in tables[1].rows[2].cells]
    assert ligne_2 == ["04-invalide.md", "Not processed", "Invalid note"]
    for code in ("aucun_constat", "note_invalide", "non_rattachee", "non_traitee"):
        assert not any(code in t for t in [c.text for row in tables[1].rows for c in row.cells])


def test_brouillon_marque_sur_les_seules_fiches_non_validees(tmp_path: Path):
    chemin = tmp_path / "brouillon.docx"
    options = OptionsExport(mode="brouillon", non_validees=frozenset({"02-notes.md"}))
    ExportDocx(tmp_path / "absent").exporter(_rapport(avec_ecartees=False), chemin, options)

    document = Document(str(chemin))
    titres = _titres(document)
    assert "1. Injection SQL sur le formulaire de connexion" in titres
    assert "2. Vecteur CVSS invalide — Non validée" in titres

    textes = _textes(document)
    assert "À revoir" in textes
    assert "vecteur_invalide : CVSS:3.1/AV:N/AC:X/PR:N/UI:N/S:U/C:N/I:N/A:N" in textes

    ligne_a_revoir = [cellule.text for cellule in document.tables[0].rows[2].cells]
    assert ligne_a_revoir[-1] == "À revoir"


def test_final_aucune_marque_validee_sans_motifs(tmp_path: Path):
    chemin = tmp_path / "rapport-final.docx"
    options = OptionsExport(mode="final")
    ExportDocx(tmp_path / "absent").exporter(_rapport(avec_ecartees=False), chemin, options)

    document = Document(str(chemin))
    titres = _titres(document)
    assert "1. Injection SQL sur le formulaire de connexion" in titres
    assert "2. Vecteur CVSS invalide" in titres
    assert not any(t.endswith("Non validée") for t in titres)

    textes = _textes(document)
    assert textes.count("Validée") == 2
    assert "À revoir" not in textes
    assert not any("vecteur_invalide" in t for t in textes)

    for ligne in (document.tables[0].rows[1].cells, document.tables[0].rows[2].cells):
        assert [c.text for c in ligne][-1] == "Validée"


@pytest.mark.parametrize(
    ("langue", "marque", "motif"),
    [
        ("fr", "Non validée", "fiche non validée, exclue du rapport final"),
        ("en", "Not validated", "not validated, excluded from the final report"),
    ],
)
def test_final_annexe_non_validees_fr_en(tmp_path: Path, langue: str, marque: str, motif: str):
    chemin = tmp_path / "rapport-final.docx"
    options = OptionsExport(
        langue=langue, mode="final", non_validees=frozenset({"09-ecartee.md", "01-autre.md"})
    )
    ExportDocx(tmp_path / "absent").exporter(
        Rapport(fiches=[_fiche_ok()], ecartees=[]), chemin, options
    )

    document = Document(str(chemin))
    titres = [p.text for p in document.paragraphs]
    assert LIBELLES[langue]["annexe_titre"] in titres

    annexe = document.tables[1]
    lignes = [[c.text for c in row.cells] for row in annexe.rows[1:]]
    assert lignes == [
        ["01-autre.md", marque, motif],
        ["09-ecartee.md", marque, motif],
    ]


def test_annexe_avec_lot_interrompu(tmp_path: Path):
    entree = EntreeEcartee(
        note="10-pas-vue.md",
        statut=Statut.NON_TRAITEE,
        motif=Motif(code=CodeMotif.LOT_INTERROMPU, detail="lot interrompu"),
        candidats=[],
        constat_attendu=None,
    )
    chemin = tmp_path / "rapport.docx"
    ExportDocx(tmp_path / "absent").exporter(
        Rapport(fiches=[_fiche_ok()], ecartees=[entree]), chemin
    )

    document = Document(str(chemin))
    annexe = document.tables[1]
    ligne = [c.text for c in annexe.rows[1].cells]
    assert ligne == ["10-pas-vue.md", "non_traitee", "lot_interrompu : lot interrompu"]


def test_aucun_partiel_restant(tmp_path: Path):
    chemin = tmp_path / "rapport.docx"
    exporteur = ExportDocx(tmp_path / "absent")
    exporteur.exporter(_rapport(avec_ecartees=True), chemin)
    exporteur.exporter(_rapport(avec_ecartees=False), chemin, OptionsExport(mode="brouillon"))

    assert [p.name for p in tmp_path.iterdir()] == ["rapport.docx"]


def test_en_mots_inline_cvss_et_tracabilite(tmp_path: Path):
    chemin = tmp_path / "rapport.docx"
    ExportDocx(tmp_path / "absent").exporter(
        _rapport(avec_ecartees=False), chemin, OptionsExport(langue="en")
    )

    textes = _textes(Document(str(chemin)))
    assert (
        "Vector : CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H — Score : 9.8 — Severity : Critical"
        in textes
    )
    assert "Note : 01-injection-sql.md — Finding template : WEB-INJECTION-SQL" in textes
    assert not any("Vecteur :" in t or "Constat type :" in t or "Sévérité :" in t for t in textes)


def test_fr_mots_inline_inchanges(tmp_path: Path):
    chemin = tmp_path / "rapport.docx"
    ExportDocx(tmp_path / "absent").exporter(_rapport(avec_ecartees=False), chemin)

    textes = _textes(Document(str(chemin)))
    assert (
        "Vecteur : CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H — Score : 9.8 — Sévérité : Critique"
        in textes
    )
    assert "Note : 01-injection-sql.md — Constat type : WEB-INJECTION-SQL" in textes


def test_aucun_partiel_si_le_remplacement_echoue(tmp_path: Path, monkeypatch):
    def refus(*_args, **_kwargs):
        raise PermissionError("document ouvert dans Word")

    monkeypatch.setattr("waveaudit.adaptateurs.export_docx.os.replace", refus)
    chemin = tmp_path / "rapport.docx"
    with pytest.raises(PermissionError):
        ExportDocx(tmp_path / "absent").exporter(_rapport(avec_ecartees=False), chemin)

    assert list(tmp_path.iterdir()) == []


# [LIB, revue] Un motif de chaque code, avec le `detail` qu'émet réellement le code source
# (garde_fous, qualite_texte.motif_langue, pipeline, moteur_openai, relecture).
def _motifs_reels() -> list[Motif]:
    return [
        Motif(code=CodeMotif.REFERENCE_RETIREE, detail="CVE-2099-0001"),
        Motif(
            code=CodeMotif.VECTEUR_INVALIDE,
            detail="CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:L/I:L/A:N",
        ),
        Motif(code=CodeMotif.RUBRIQUE_COMPLETEE, detail="recommandation"),
        Motif(code=CodeMotif.NOTE_TRONQUEE, detail="corps tronqué à 500 caractères"),
        Motif(
            code=CodeMotif.REPLI_SORTIE_INVALIDE,
            detail="réponse du moteur : contenu non JSON",
        ),
        Motif(
            code=CodeMotif.REPLI_MOTEUR,
            detail="moteur injoignable ([WinError 10061] connexion refusée)",
        ),
        Motif(code=CodeMotif.NOTE_INVALIDE, detail="corps vide"),
        Motif(code=CodeMotif.AUCUN_CONSTAT, detail="score BM25 maximal nul"),
        Motif(code=CodeMotif.INDICATEUR_RETIRE, detail="203.0.113.7"),
        Motif(code=CodeMotif.REPETITION, detail="titre"),
        Motif(
            code=CodeMotif.LANGUE,
            detail=f"description : {motif_langue(_TEXTE_FRANCAIS, 'en')}",
        ),
        Motif(code=CodeMotif.LOT_INTERROMPU, detail="lot interrompu avant cette note"),
    ]


_TEXTE_FRANCAIS = (
    "Le paramètre de la page est renvoyé dans la réponse et il est exécuté par le navigateur."
)


def _rapport_tous_motifs() -> Rapport:
    fiche = _fiche_a_revoir().model_copy(
        update={
            "titre": "Reflected XSS",
            "description": "The parameter is echoed back.",
            "impact": "Session theft.",
            "recommandation": "Encode the output.",
            "motifs": _motifs_reels(),
        }
    )
    ecartees = [
        EntreeEcartee(
            note=f"{rang:02d}-note.md",
            statut=Statut.NON_TRAITEE if rang % 2 else Statut.NON_RATTACHEE,
            motif=motif,
            candidats=[],
            constat_attendu=None,
        )
        for rang, motif in enumerate(_motifs_reels(), start=10)
    ]
    return Rapport(fiches=[fiche], ecartees=ecartees)


def _tous_les_textes(document: Document) -> list[str]:
    cellules = [c.text for t in document.tables for row in t.rows for c in row.cells]
    return _textes(document) + cellules


# Mots français des `detail` réels ci-dessus et des statuts : aucun n'atteint un Word en.
_MOTS_FRANCAIS = re.compile(
    r"\b(corps|tronqué|caractères|réponse|moteur|contenu|injoignable|connexion|refusée|vide|"
    r"maximal|nul|lot|interrompu|avant|cette|titre|recommandation|écriture|latine|mots-outils|"
    r"traitée|rattachée|aucun|candidat|seuil)\b",
    re.IGNORECASE,
)
_SNAKE_CASE = re.compile(r"\b[a-z0-9]+(?:_[a-z0-9]+)+\b")


@pytest.mark.parametrize("mode", ["lot", "brouillon", "final"])
def test_en_tous_les_codes_sans_francais_ni_code_brut(tmp_path: Path, mode: str):
    chemin = tmp_path / "rapport.docx"
    options = OptionsExport(langue="en", mode=mode)
    ExportDocx(tmp_path / "absent").exporter(_rapport_tous_motifs(), chemin, options)

    document = Document(str(chemin))
    textes = _tous_les_textes(document)
    for texte in textes:
        assert _MOTS_FRANCAIS.search(texte) is None, texte
        assert _SNAKE_CASE.search(texte) is None, texte
    annexe = document.tables[1]
    motifs = [row.cells[2].text for row in annexe.rows[1:13]]
    assert motifs == [
        "Reference removed : CVE-2099-0001",
        "Invalid CVSS vector : CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:L/I:L/A:N",
        "Section completed : Recommendation",
        "Note truncated : cut at 500 characters",
        "Fallback — invalid output",
        "Engine fallback",
        "Invalid note",
        "No matching finding",
        "Indicator removed : 203.0.113.7",
        "Repetitive section : Title",
        "Language mismatch : Description, fr function words",
        "Batch interrupted",
    ]
    assert [row.cells[1].text for row in annexe.rows[1:3]] == ["Not matched", "Not processed"]
    if mode != "final":
        # Corps de la fiche à revoir : mêmes rendus que dans l'annexe.
        assert "Section completed : Recommendation" in textes
        assert "Engine fallback" in textes


@pytest.mark.parametrize("mode", ["lot", "brouillon", "final"])
def test_fr_rendu_v010_code_et_detail_bruts(tmp_path: Path, mode: str):
    """§17 et §27 : le Word fr reste celui de la v0.1.0 (statut brut, « code : detail »)."""
    chemin = tmp_path / "rapport.docx"
    ExportDocx(tmp_path / "absent").exporter(
        _rapport_tous_motifs(), chemin, OptionsExport(mode=mode)
    )

    document = Document(str(chemin))
    annexe = document.tables[1]
    lignes = [[c.text for c in row.cells] for row in annexe.rows[1:13]]
    assert lignes[0] == ["10-note.md", "non_rattachee", "reference_retiree : CVE-2099-0001"]
    assert lignes[3] == [
        "13-note.md",
        "non_traitee",
        "note_tronquee : corps tronqué à 500 caractères",
    ]
    assert lignes[5][2] == "repli_moteur : moteur injoignable ([WinError 10061] connexion refusée)"
    assert lignes[10][2] == "langue : description : mots-outils fr"
    assert lignes[11][2] == "lot_interrompu : lot interrompu avant cette note"
    if mode != "final":
        textes = _textes(document)
        assert "repetition : titre" in textes
        assert "aucun_constat : score BM25 maximal nul" in textes
        assert len([t for t in textes if " : " in t and t.split(" : ")[0] in set(CodeMotif)]) == 12


def test_note_tronquee_format_lie_au_pipeline():
    """Le rendu anglais extrait le nombre du `detail` que construit réellement le pipeline."""
    source = inspect.getsource(pipeline)
    assert 'detail=f"corps tronqué à {moteur.taille_max_note} caractères"' in source
    motif = Motif(code=CodeMotif.NOTE_TRONQUEE, detail="corps tronqué à 80 caractères")
    assert _motif_affiche(motif, "en") == "Note truncated : cut at 80 characters"


def test_langue_suffixes_lies_a_qualite_texte():
    assert motif_langue("Внедрение SQL", "fr") == "écriture non latine"
    assert motif_langue(_TEXTE_FRANCAIS, "en") == "mots-outils fr"
    non_latin = Motif(code=CodeMotif.LANGUE, detail="impact : écriture non latine")
    assert _motif_affiche(non_latin, "en") == "Language mismatch : Impact, non-Latin script"


def test_en_format_de_detail_inconnu_jamais_recopie():
    def rendu(code: CodeMotif, detail: str) -> str:
        return _motif_affiche(Motif(code=code, detail=detail), "en")

    assert rendu(CodeMotif.LANGUE, "impact : autre chose") == "Language mismatch : Impact"
    assert rendu(CodeMotif.LANGUE, "xx : autre chose") == "Language mismatch"
    assert rendu(CodeMotif.NOTE_TRONQUEE, "autre chose") == "Note truncated"
    assert rendu(CodeMotif.REPETITION, "autre chose") == "Repetitive section"


def test_completude_classement_des_codes_et_statuts():
    assert set(MOTIFS_EN) == set(CodeMotif)
    assert set(DETAILS_EN) == set(CodeMotif)
    assert set(RUBRIQUES_EN) == set(RUBRIQUES)
    statuts = typing.get_args(EntreeEcartee.model_fields["statut"].annotation)
    assert set(STATUTS_ECARTES_EN) == set(statuts)
