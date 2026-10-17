"""Tests de l'export Word du rapport d'incident (CAP-6, contrats §21, formats-incident.md)."""

from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from docx import Document

from tests.incident_outils import detection, evenement, recit
from waveaudit.adaptateurs.export_docx import BANDEAU
from waveaudit.adaptateurs.export_incident_docx import ExporteurIncidentDocx
from waveaudit.domaine.erreurs import ConfigurationInvalide
from waveaudit.domaine.incident import Indicateur, LigneIgnoree, RapportIncident
from waveaudit.domaine.modeles import CodeMotif, Motif, Statut

FUSEAU = timezone(timedelta(hours=2))


def _rapport_avec_fait() -> RapportIncident:
    evenements = [
        evenement(ligne=1, horodatage=datetime(2026, 10, 14, 8, 2, 10, tzinfo=FUSEAU)),
        evenement(ligne=2, horodatage=datetime(2026, 10, 14, 8, 2, 20, tzinfo=FUSEAU)),
        evenement(ligne=3, horodatage=datetime(2026, 10, 14, 8, 2, 30, tzinfo=FUSEAU)),
    ]
    det = detection(evenements=evenements)
    r_ok = recit(detection=det)
    r_a_revoir = recit(
        detection=det,
        regle_titre="Force brute SSH",
        severite="critique",
        statut=Statut.A_REVOIR,
        motifs=[Motif(code=CodeMotif.INDICATEUR_RETIRE, detail="203.0.113.254")],
    )
    return RapportIncident(
        debut=evenements[0].horodatage,
        fin=evenements[-1].horodatage,
        recits=[r_ok, r_a_revoir],
        indicateurs=[
            Indicateur(type="ip", valeur="203.0.113.45", origines=["auth.log:1", "auth.log:2"])
        ],
        actifs=["BASTION01"],
        recommandations=["Bloquer l'adresse source et imposer l'authentification par clé."],
        sources={"auth.log": 10, "security.json": 5},
        ignorees=[LigneIgnoree(fichier="auth.log", ligne=57, raison="illisible")],
    )


def _rapport_vide() -> RapportIncident:
    return RapportIncident(
        debut=None,
        fin=None,
        recits=[],
        indicateurs=[],
        actifs=[],
        recommandations=[],
        sources={"auth.log": 4},
        ignorees=[],
    )


def _titres(document: Document) -> list[str]:
    return [p.text for p in document.paragraphs if p.style.name.startswith(("Heading", "Title"))]


def _textes(document: Document) -> list[str]:
    return [p.text for p in document.paragraphs]


def test_rapport_avec_fait_contient_les_sections_dans_lordre(tmp_path: Path):
    chemin = tmp_path / "sortie" / "rapport-incident.docx"
    ExporteurIncidentDocx(tmp_path / "theme-prive-absent").exporter(_rapport_avec_fait(), chemin)

    assert chemin.is_file()
    document = Document(str(chemin))

    for section in document.sections:
        assert section.header.paragraphs[0].text == BANDEAU
        assert section.footer.paragraphs[0].text == BANDEAU

    titres = _titres(document)
    ordre_attendu = [
        "Rapport d'incident",
        "Synthèse",
        "Chronologie",
        "Faits détectés",
        "Indicateurs de compromission",
        "Actifs touchés",
        "Recommandations consolidées",
        "Annexe : sources analysées",
    ]
    indices = [titres.index(t) for t in ordre_attendu]
    assert indices == sorted(indices)

    textes = _textes(document)
    assert any("Période : 2026-10-14T08:02:10+02:00" in t for t in textes)
    assert any("Hôtes touchés : BASTION01" in t for t in textes)

    tables = document.tables
    synthese = tables[0]
    lignes = [[c.text for c in row.cells] for row in synthese.rows]
    assert lignes[0] == ["Sévérité", "Nombre"]
    assert lignes[1] == ["Critique", "1"]
    assert lignes[2] == ["Élevée", "1"]
    assert lignes[3] == ["Moyenne", "0"]
    assert lignes[4] == ["Faible", "0"]

    chronologie = tables[1]
    assert [c.text for c in chronologie.rows[0].cells] == ["Horodatage", "Hôte", "Fait", "Sévérité"]
    ligne_1 = [c.text for c in chronologie.rows[1].cells]
    assert ligne_1 == ["2026-10-14T08:02:10+02:00", "BASTION01", "Force brute SSH", "Élevée"]
    ligne_2 = [c.text for c in chronologie.rows[2].cells]
    assert ligne_2[-1] == "Critique"

    indicateurs = tables[2]
    assert [c.text for c in indicateurs.rows[0].cells] == ["Type", "Valeur", "Origines"]
    assert [c.text for c in indicateurs.rows[1].cells] == [
        "ip",
        "203.0.113.45",
        "auth.log:1, auth.log:2",
    ]

    sources = tables[3]
    assert [c.text for c in sources.rows[0].cells] == [
        "Fichier",
        "Enregistrements examinés",
        "Ignorés",
    ]
    lignes_sources = {row.cells[0].text: [c.text for c in row.cells] for row in sources.rows[1:]}
    assert lignes_sources["auth.log"] == ["auth.log", "10", "1"]
    assert lignes_sources["security.json"] == ["security.json", "5", "0"]
    assert any("security.json : ligne = rang d'objet" in t for t in textes)

    assert any("Bloquer l'adresse source" in t for t in textes)
    assert "À revoir" in textes
    assert any("indicateur_retire : 203.0.113.254" in t for t in textes)


def test_preuves_limitees_a_max_preuves_avec_reste(tmp_path: Path):
    evenements = [
        evenement(ligne=i, horodatage=datetime(2026, 10, 14, 8, i, 0, tzinfo=FUSEAU))
        for i in range(2, 10)
    ]
    det = detection(evenements=evenements)
    rapport = RapportIncident(
        debut=evenements[0].horodatage,
        fin=evenements[-1].horodatage,
        recits=[recit(detection=det)],
        indicateurs=[],
        actifs=["BASTION01"],
        recommandations=[],
        sources={"auth.log": 8},
        ignorees=[],
    )
    chemin = tmp_path / "rapport.docx"
    ExporteurIncidentDocx(tmp_path / "absent", max_preuves=3).exporter(rapport, chemin)

    textes = _textes(Document(str(chemin)))
    assert sum(t.startswith("auth.log:") for t in textes) == 3
    assert "+ 5 autres" in textes


def test_sans_detection_mentionne_aucun_fait_detecte(tmp_path: Path):
    chemin = tmp_path / "rapport.docx"
    ExporteurIncidentDocx(tmp_path / "absent").exporter(_rapport_vide(), chemin)

    document = Document(str(chemin))
    textes = _textes(document)
    assert textes.count("Aucun fait détecté") == 5
    tables = document.tables
    assert len(tables) == 2  # synthèse (zéros) + annexe des sources
    synthese = [[c.text for c in row.cells] for row in tables[0].rows]
    assert synthese[1:] == [["Critique", "0"], ["Élevée", "0"], ["Moyenne", "0"], ["Faible", "0"]]


def test_caracteres_interdits_en_xml_retires(tmp_path: Path):
    r = recit(titre="T\x00itre\x0b", recit="Un récit\x1fpropre.")
    rapport = RapportIncident(
        debut=r.detection.debut,
        fin=r.detection.fin,
        recits=[r],
        indicateurs=[],
        actifs=[],
        recommandations=[],
        sources={},
        ignorees=[],
    )
    chemin = tmp_path / "rapport.docx"
    ExporteurIncidentDocx(tmp_path / "absent").exporter(rapport, chemin)

    textes = _textes(Document(str(chemin)))
    assert any(t.startswith("1. Titre (") for t in textes)
    assert "Un récitpropre." in textes
    assert [p.name for p in tmp_path.iterdir()] == ["rapport.docx"]


def test_theme_prive_avec_modele_et_theme_toml(tmp_path: Path):
    theme_prive = tmp_path / "theme-prive"
    theme_prive.mkdir()
    Document().save(str(theme_prive / "modele.docx"))
    (theme_prive / "theme.toml").write_text('couleur_primaire = "#123456"\n', encoding="utf-8")

    chemin = tmp_path / "rapport.docx"
    ExporteurIncidentDocx(theme_prive).exporter(_rapport_vide(), chemin)

    document = Document(str(chemin))
    assert str(document.styles["Heading 1"].font.color.rgb) == "123456"
    assert "Rapport d'incident" in _titres(document)


@pytest.mark.parametrize(
    "contenu",
    ['couleur_primaire = "rouge"', "x ="],
)
def test_theme_toml_invalide_leve_configuration_invalide(tmp_path: Path, contenu: str):
    theme_prive = tmp_path / "theme-prive"
    theme_prive.mkdir()
    (theme_prive / "theme.toml").write_text(contenu, encoding="utf-8")

    with pytest.raises(ConfigurationInvalide, match="theme.toml"):
        ExporteurIncidentDocx(theme_prive)


def test_preuves_message_normalise_et_tronque_sans_reste(tmp_path: Path):
    message = "Failed  password\n\tfor " + "x" * 400
    evenements = [evenement(ligne=i, message=message) for i in (1, 2)]
    rapport = _rapport_vide().model_copy(
        update={"recits": [recit(detection=detection(evenements=evenements))]}
    )
    chemin = tmp_path / "rapport.docx"
    ExporteurIncidentDocx(tmp_path / "absent", max_preuves=2).exporter(rapport, chemin)

    textes = _textes(Document(str(chemin)))
    preuves = [t for t in textes if t.startswith("auth.log:")]
    attendu = ("Failed password for " + "x" * 400)[:300]
    assert preuves == [f"auth.log:1 — {attendu}", f"auth.log:2 — {attendu}"]
    assert not any(t.startswith("+ ") for t in textes)


def test_bandeau_sur_premiere_page_et_pages_paires_du_modele(tmp_path: Path):
    theme_prive = tmp_path / "theme-prive"
    theme_prive.mkdir()
    modele = Document()
    modele.settings.odd_and_even_pages_header_footer = True
    modele.sections[0].different_first_page_header_footer = True
    modele.save(str(theme_prive / "modele.docx"))

    chemin = tmp_path / "rapport.docx"
    ExporteurIncidentDocx(theme_prive).exporter(_rapport_vide(), chemin)

    section = Document(str(chemin)).sections[0]
    for conteneur in (
        section.first_page_header,
        section.first_page_footer,
        section.even_page_header,
        section.even_page_footer,
    ):
        assert BANDEAU in [p.text for p in conteneur.paragraphs]


@pytest.mark.parametrize("max_preuves", [0, -1])
def test_max_preuves_invalide_leve_configuration_invalide(tmp_path: Path, max_preuves: int):
    with pytest.raises(ConfigurationInvalide, match="max_preuves"):
        ExporteurIncidentDocx(tmp_path / "absent", max_preuves=max_preuves)


def test_modele_docx_illisible_leve_configuration_invalide(tmp_path: Path):
    theme_prive = tmp_path / "theme-prive"
    theme_prive.mkdir()
    (theme_prive / "modele.docx").write_bytes(b"pas un docx")

    with pytest.raises(ConfigurationInvalide, match="modele.docx"):
        ExporteurIncidentDocx(theme_prive)
